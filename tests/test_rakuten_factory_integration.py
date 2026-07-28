"""Integration tests for Rakuten marketplace factory and profit pipeline."""

import json
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from config.constants import MARKETPLACE_RAKUTEN, SHEET_DOMESTIC_LISTINGS
from excel.exporter import ExcelExporter
from excel.template import MARKETPLACE_LISTING_COLUMNS
from marketplace.base_marketplace import BaseMarketplace
from marketplace.marketplace_factory import create_marketplace, get_all_marketplaces
from marketplace.rakuten_client import FakeRakutenClient
from marketplace.rakuten_marketplace import RakutenMarketplace
from marketplace.rakuten_settings import RakutenConfig
from models.marketplace_listing import MarketplaceListing
from models.product import Product
from price_compare.marketplace_profit_service import calculate_profit_from_search_results
from price_compare.profit_calculator import ProfitCalculator
from price_compare.profit_config import ProfitConfig

FIXTURES = Path(__file__).parent / "fixtures"


def _rakuten_config(**overrides) -> RakutenConfig:
    defaults = dict(
        application_id="dummy-app-id",
        access_key="dummy-access-key",
        affiliate_id="",
        base_url="https://openapi.rakuten.co.jp/ichibams/api/IchibaItem/Search/20260701",
        timeout_seconds=10,
        max_retries=0,
        hits=20,
        sort="standard",
        enabled=True,
        demo_enabled=True,
    )
    defaults.update(overrides)
    return RakutenConfig(**defaults)


def _rakuten_marketplace(name: str = "rakuten_search_multiple.json") -> RakutenMarketplace:
    payload = json.loads((FIXTURES / name).read_text(encoding="utf-8"))
    return RakutenMarketplace(client=FakeRakutenClient(payload), config=_rakuten_config())


def test_create_rakuten_marketplace() -> None:
    marketplace = create_marketplace("rakuten", rakuten_settings=_rakuten_config())
    assert isinstance(marketplace, RakutenMarketplace)
    assert isinstance(marketplace, BaseMarketplace)


def test_rakuten_case_insensitive() -> None:
    marketplace = create_marketplace("RAKUTEN", rakuten_settings=_rakuten_config())
    assert isinstance(marketplace, RakutenMarketplace)


def test_other_marketplaces_preserved() -> None:
    assert create_marketplace("local").marketplace_name == "local"
    assert create_marketplace("yahoo").marketplace_name == "yahoo"
    assert create_marketplace("amazon_jp").marketplace_name == "amazon_jp"


def test_unknown_marketplace_raises() -> None:
    with pytest.raises(ValueError, match="Unsupported marketplace"):
        create_marketplace("ebay")


def test_get_all_marketplaces_includes_rakuten() -> None:
    marketplaces = get_all_marketplaces(rakuten_settings=_rakuten_config())
    names = [item.marketplace_name for item in marketplaces]
    assert MARKETPLACE_RAKUTEN in names
    assert len(marketplaces) == 5


def test_profit_with_rakuten_price() -> None:
    marketplace = _rakuten_marketplace("rakuten_search_normal.json")
    product = Product(name="Gucci Wallet", brand="GUCCI", model="456126", price=600.0, currency="USD", exchange_rate=150.0)
    result = marketplace.search(product)
    calculator = ProfitCalculator(ProfitConfig(marketplace_fee_rate=Decimal("0")))
    price_result = calculate_profit_from_search_results([result], calculator)[0]
    assert price_result.domestic_sale_price_jpy == Decimal("128000")
    assert price_result.marketplace_fee_jpy == Decimal("0")


def test_profit_free_shipping() -> None:
    marketplace = _rakuten_marketplace("rakuten_search_normal.json")
    product = Product(name="Gucci Wallet", brand="GUCCI", model="456126", price=600.0, currency="USD", exchange_rate=150.0)
    result = marketplace.search(product)
    assert result.selected_listing is not None
    assert result.selected_listing.shipping_jpy == Decimal("0")
    assert result.selected_price_jpy == Decimal("128000")


def test_profit_unknown_shipping() -> None:
    marketplace = _rakuten_marketplace("rakuten_search_multiple.json")
    product = Product(name="Gucci Bag", brand="GUCCI", model="GG-MARMONT", price=600.0, currency="USD", exchange_rate=150.0)
    result = marketplace.search(product)
    selected = result.selected_listing
    assert selected is not None
    premium = next(i for i in result.valid_listings if i.listing_id == "premium-shop:789012")
    assert premium.shipping_unknown is True
    assert result.selected_price_jpy == Decimal("198000")


def test_points_not_deducted_from_profit() -> None:
    marketplace = _rakuten_marketplace("rakuten_search_normal.json")
    product = Product(name="Gucci Wallet", brand="GUCCI", model="456126", price=600.0, currency="USD", exchange_rate=150.0)
    result = marketplace.search(product)
    assert result.selected_listing is not None
    assert result.selected_listing.point_rate == Decimal("1")
    calculator = ProfitCalculator()
    price_result = calculate_profit_from_search_results([result], calculator)[0]
    assert price_result.domestic_sale_price_jpy == Decimal("128000")


def test_excel_rakuten_rows(tmp_path: Path) -> None:
    marketplace = _rakuten_marketplace("rakuten_search_multiple.json")
    product = Product(name="Gucci Bag", brand="GUCCI", model="GG-MARMONT", price=600.0, currency="USD", exchange_rate=150.0)
    result = marketplace.search(product)
    calculator = ProfitCalculator()
    price_results = calculate_profit_from_search_results([result], calculator)
    exporter = ExcelExporter(output_dir=tmp_path, filename="rakuten.xlsx")
    exporter.export_phase3_workbook([product], result.listings, price_results)

    workbook = openpyxl.load_workbook(exporter.output_path)
    sheet = workbook[SHEET_DOMESTIC_LISTINGS]
    marketplace_col = MARKETPLACE_LISTING_COLUMNS.index("marketplace_name") + 1
    values = [sheet.cell(row=row, column=marketplace_col).value for row in range(2, sheet.max_row + 1)]
    assert MARKETPLACE_RAKUTEN in values

    price_col = MARKETPLACE_LISTING_COLUMNS.index("price_jpy") + 1
    assert isinstance(sheet.cell(row=2, column=price_col).value, (int, float))

    point_rate_col = MARKETPLACE_LISTING_COLUMNS.index("point_rate") + 1
    point_value = sheet.cell(row=2, column=point_rate_col).value
    assert point_value is None or isinstance(point_value, (int, float))
    assert point_value != "None"


def test_excel_empty_rakuten(tmp_path: Path) -> None:
    marketplace = _rakuten_marketplace("rakuten_search_empty.json")
    product = Product(name="Unknown", brand="Demo", price=100.0)
    result = marketplace.search(product)
    exporter = ExcelExporter(output_dir=tmp_path, filename="rakuten_empty.xlsx")
    exporter.export_phase3_workbook([product], result.listings, [])
    workbook = openpyxl.load_workbook(exporter.output_path)
    assert SHEET_DOMESTIC_LISTINGS in workbook.sheetnames


def test_excel_yahoo_amazon_preserved(tmp_path: Path) -> None:
    yahoo = MarketplaceListing(marketplace_name="yahoo", listing_id="Y-1", title="Yahoo", price_jpy=Decimal("10000"), shipping_jpy=Decimal("0"), listing_url="https://example.com/y")
    amazon = MarketplaceListing(marketplace_name="amazon_jp", listing_id="A-1", title="Amazon", price_jpy=Decimal("20000"), shipping_jpy=Decimal("0"), listing_url="https://example.com/a")
    rakuten = _rakuten_marketplace("rakuten_search_normal.json").search(Product(name="Gucci Wallet", brand="GUCCI", model="456126", price=100.0)).listings[0]
    exporter = ExcelExporter(output_dir=tmp_path, filename="mixed.xlsx")
    exporter.export_phase3_workbook([], [yahoo, amazon, rakuten], [])
    workbook = openpyxl.load_workbook(exporter.output_path)
    assert workbook[SHEET_DOMESTIC_LISTINGS].max_row == 4
