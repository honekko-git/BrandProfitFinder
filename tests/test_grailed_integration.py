"""Integration tests for Grailed marketplace."""

import json
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from config.constants import MARKETPLACE_GRAILED, SHEET_DOMESTIC_LISTINGS
from excel.exporter import ExcelExporter
from excel.template import MARKETPLACE_LISTING_COLUMNS
from marketplace.grailed_client import FakeGrailedClient
from marketplace.grailed_exceptions import GrailedConfigurationError
from marketplace.grailed_marketplace import GrailedMarketplace, create_grailed_marketplace
from marketplace.grailed_settings import GrailedSettings
from marketplace.marketplace_factory import create_marketplace
from models.product import Product
from price_compare.marketplace_profit_service import calculate_profit_from_search_results
from price_compare.profit_calculator import ProfitCalculator

FIXTURES = Path(__file__).parent / "fixtures"


def _settings(**overrides) -> GrailedSettings:
    defaults = dict(
        enabled=True,
        timeout_seconds=10,
        page_size=20,
        max_pages=2,
        default_currency="JPY",
        demo_fixture_path="grailed_search_normal.json",
        allow_unknown_currency=False,
        include_unavailable=False,
        include_sold=False,
        include_reserved=False,
        include_offer_enabled_only=False,
        include_discounted_only=False,
        require_verified_seller=False,
    )
    defaults.update(overrides)
    return GrailedSettings(**defaults)


def _marketplace(name: str = "grailed_search_normal.json", **settings) -> GrailedMarketplace:
    payload = json.loads((FIXTURES / name).read_text(encoding="utf-8"))
    return GrailedMarketplace(
        client=FakeGrailedClient(payload),
        settings=_settings(**settings),
    )


def test_marketplace_search() -> None:
    product = Product(
        name="Ramones",
        brand="RICK OWENS",
        model="DRKSHDW-RAMONES",
        price=600.0,
        currency="USD",
        exchange_rate=150.0,
    )
    result = _marketplace().search(product)
    assert result.marketplace_name == MARKETPLACE_GRAILED
    assert len(result.listings) == 1
    assert result.selected_price_jpy == Decimal("98000")


def test_empty_query() -> None:
    result = _marketplace().search(Product(name="", brand="", model=""), query="  ")
    assert result.status == "error"


def test_pagination_two_pages() -> None:
    p1 = json.loads((FIXTURES / "grailed_search_page_1.json").read_text(encoding="utf-8"))
    p2 = json.loads((FIXTURES / "grailed_search_page_2.json").read_text(encoding="utf-8"))
    mp = GrailedMarketplace(
        client=FakeGrailedClient(pages={1: p1, 2: p2}),
        settings=_settings(max_pages=2),
    )
    result = mp.search(Product(name="Sneakers", brand="RICK OWENS", model="DRKSHDW-RAMONES", price=100.0))
    assert len(result.listings) == 2


def test_factory_aliases() -> None:
    payload = json.loads((FIXTURES / "grailed_search_normal.json").read_text(encoding="utf-8"))
    client = FakeGrailedClient(payload)
    settings = _settings()
    for alias in ("grailed_market", "grailed-market", "gr"):
        mp = create_marketplace(alias, grailed_client=client, grailed_settings=settings)
        assert isinstance(mp, GrailedMarketplace)


def test_factory_no_fake_without_client() -> None:
    with pytest.raises(GrailedConfigurationError):
        create_marketplace("grailed", grailed_settings=_settings())


def test_profit_metadata() -> None:
    product = Product(
        name="Ramones",
        brand="RICK OWENS",
        model="DRKSHDW-RAMONES",
        price=600.0,
        currency="USD",
        exchange_rate=150.0,
    )
    result = _marketplace().search(product)
    price_result = calculate_profit_from_search_results([result], ProfitCalculator())[0]
    assert price_result.domestic_sale_price_jpy == Decimal("98000")
    assert price_result.metadata.get("discount_applied") is False
    assert price_result.metadata.get("offer_applied") is False
    assert price_result.metadata.get("adjustment_applied") is False
    assert price_result.metadata.get("offer_enabled") is True
    assert price_result.metadata.get("seller_transactions") == 124


def test_excel_rows(tmp_path: Path) -> None:
    product = Product(
        name="Ramones",
        brand="RICK OWENS",
        model="DRKSHDW-RAMONES",
        price=600.0,
        currency="USD",
        exchange_rate=150.0,
    )
    result = _marketplace("grailed_search_multiple.json").search(product)
    price_results = calculate_profit_from_search_results([result], ProfitCalculator())
    output_path = ExcelExporter(output_dir=tmp_path, filename="grailed.xlsx").export_phase3_workbook(
        [product], result.listings, price_results
    )
    sheet = openpyxl.load_workbook(output_path)[SHEET_DOMESTIC_LISTINGS]
    col = MARKETPLACE_LISTING_COLUMNS.index("marketplace_name") + 1
    assert MARKETPLACE_GRAILED in [sheet.cell(row=r, column=col).value for r in range(2, sheet.max_row + 1)]
    offer_col = MARKETPLACE_LISTING_COLUMNS.index("source_minimum_offer") + 1
    assert sheet.cell(row=2, column=offer_col).value is not None
