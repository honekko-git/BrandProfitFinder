"""Integration tests for Chrono24 marketplace."""

import json
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from config.constants import MARKETPLACE_CHRONO24, SHEET_DOMESTIC_LISTINGS
from excel.exporter import ExcelExporter
from excel.template import MARKETPLACE_LISTING_COLUMNS
from marketplace.chrono24_client import FakeChrono24Client
from marketplace.chrono24_exceptions import Chrono24ConfigurationError
from marketplace.chrono24_marketplace import Chrono24Marketplace, create_chrono24_marketplace
from marketplace.chrono24_settings import Chrono24Settings
from marketplace.marketplace_factory import create_marketplace
from models.product import Product
from price_compare.marketplace_profit_service import calculate_profit_from_search_results
from price_compare.profit_calculator import ProfitCalculator

FIXTURES = Path(__file__).parent / "fixtures"


def _settings(**overrides) -> Chrono24Settings:
    defaults = dict(
        enabled=True,
        timeout_seconds=10,
        page_size=20,
        max_pages=2,
        default_currency="JPY",
        demo_fixture_path="chrono24_search_normal.json",
        allow_unknown_currency=False,
        include_unavailable=False,
        include_sold=False,
        include_reserved=False,
        include_negotiable_only=False,
        include_discounted_only=False,
        require_verified_seller=False,
        require_trusted_seller=False,
        include_private_sellers=True,
        include_professional_dealers=True,
    )
    defaults.update(overrides)
    return Chrono24Settings(**defaults)


def _marketplace(name: str = "chrono24_search_normal.json", **settings) -> Chrono24Marketplace:
    payload = json.loads((FIXTURES / name).read_text(encoding="utf-8"))
    return Chrono24Marketplace(
        client=FakeChrono24Client(payload),
        settings=_settings(**settings),
    )


def test_marketplace_search() -> None:
    product = Product(
        name="Submariner",
        brand="ROLEX",
        model="126610LN",
        price=12000.0,
        currency="USD",
        exchange_rate=150.0,
    )
    result = _marketplace().search(product)
    assert result.marketplace_name == MARKETPLACE_CHRONO24
    assert len(result.listings) == 1
    assert result.selected_price_jpy == Decimal("1980000")


def test_factory_aliases() -> None:
    payload = json.loads((FIXTURES / "chrono24_search_normal.json").read_text(encoding="utf-8"))
    client = FakeChrono24Client(payload)
    settings = _settings()
    for alias in ("chrono_24", "chrono-24", "c24"):
        mp = create_marketplace(alias, chrono24_client=client, chrono24_settings=settings)
        assert isinstance(mp, Chrono24Marketplace)


def test_factory_no_fake_without_client() -> None:
    with pytest.raises(Chrono24ConfigurationError):
        create_marketplace("chrono24", chrono24_settings=_settings())


def test_profit_metadata() -> None:
    product = Product(name="Submariner", brand="ROLEX", model="126610LN", price=12000.0, currency="USD", exchange_rate=150.0)
    result = _marketplace().search(product)
    price_result = calculate_profit_from_search_results([result], ProfitCalculator())[0]
    assert price_result.domestic_sale_price_jpy == Decimal("1980000")
    assert price_result.metadata.get("negotiation_applied") is False
    assert price_result.metadata.get("reference_number") == "126610LN"
    assert price_result.metadata.get("case_diameter_mm") == 41.0


def test_excel_rows(tmp_path: Path) -> None:
    product = Product(name="Submariner", brand="ROLEX", model="126610LN", price=12000.0, currency="USD", exchange_rate=150.0)
    result = _marketplace("chrono24_search_multiple.json").search(product)
    price_results = calculate_profit_from_search_results([result], ProfitCalculator())
    output_path = ExcelExporter(output_dir=tmp_path, filename="c24.xlsx").export_phase3_workbook(
        [product], result.listings, price_results
    )
    sheet = openpyxl.load_workbook(output_path)[SHEET_DOMESTIC_LISTINGS]
    col = MARKETPLACE_LISTING_COLUMNS.index("source_reference_number") + 1
    assert "126610LN" in [sheet.cell(row=r, column=col).value for r in range(2, sheet.max_row + 1)]
