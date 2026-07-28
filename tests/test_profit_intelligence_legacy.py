"""Legacy compatibility tests when Profit Intelligence is disabled."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest

import main
from excel.exporter import ExcelExporter, SHEET_PRICE_RESULTS
from excel.template import PRICE_RESULT_COLUMNS, price_result_columns
from models.product import Product
from price_compare.profit_calculator import ProfitCalculator
from price_compare.profit_config import ProfitConfig
from price_compare.ranking_engine import RankingEngine, RankingSortKey
from profit_intelligence.service import ProfitIntelligenceService


def _phase3_calculator() -> ProfitCalculator:
    return ProfitCalculator(
        ProfitConfig(
            international_shipping_jpy=Decimal("2500"),
            customs_duty_rate=Decimal("0.08"),
            import_tax_rate=Decimal("0.10"),
            domestic_shipping_jpy=Decimal("800"),
            marketplace_fee_rate=Decimal("0.12"),
            other_costs_jpy=Decimal("500"),
        )
    )


def test_rank_helper_matches_legacy_ordering_when_disabled() -> None:
    products = main.build_phase3_products()
    listings_map = main.build_phase3_listings()
    marketplace = main.create_marketplace(
        "local",
        listings_by_product_key=listings_map,
        selection_strategy=main.PriceSelectionStrategy.HIGHEST,
    )
    calculator = _phase3_calculator()
    search_results = [marketplace.search(product) for product in products]
    results = main.calculate_profit_from_search_results(search_results, calculator)
    legacy = RankingEngine(calculator.config).rank(
        results,
        sort_key=RankingSortKey.PROFIT,
        descending=True,
    )
    via_helper = main._rank_phase3_results(results, search_results, calculator, False)
    assert [item.profit_jpy for item in legacy] == [item.profit_jpy for item in via_helper]
    assert all(item.profit_intelligence is None for item in via_helper)


def test_scoring_disabled_does_not_invoke_service() -> None:
    products = main.build_phase3_products()
    listings_map = main.build_phase3_listings()
    marketplace = main.create_marketplace(
        "local",
        listings_by_product_key=listings_map,
        selection_strategy=main.PriceSelectionStrategy.HIGHEST,
    )
    calculator = _phase3_calculator()
    search_results = [marketplace.search(product) for product in products]
    results = main.calculate_profit_from_search_results(search_results, calculator)
    with patch.object(ProfitIntelligenceService, "score_results") as score_mock:
        main._rank_phase3_results(results, search_results, calculator, False)
        score_mock.assert_not_called()


def test_legacy_excel_headers_unchanged(tmp_path: Path) -> None:
    product = Product(name="Bag", brand="Demo", price=100.0, currency="USD", exchange_rate=150.0)
    result = ProfitCalculator().calculate(product, Decimal("50000"))
    exporter = ExcelExporter(output_dir=tmp_path, filename="legacy.xlsx")
    exporter.export_phase3_workbook([product], [], [result])
    workbook = openpyxl.load_workbook(exporter.output_path)
    sheet = workbook[SHEET_PRICE_RESULTS]
    headers = [sheet.cell(row=1, column=i + 1).value for i in range(len(PRICE_RESULT_COLUMNS))]
    assert headers == PRICE_RESULT_COLUMNS
    assert price_result_columns([result]) == PRICE_RESULT_COLUMNS


def test_legacy_run_does_not_attach_intelligence(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(main, "OUTPUT_DIR", tmp_path)
    with patch("utils.http.fetch_url"), patch("utils.http.HttpClient"):
        main.run(profit_intelligence=False)
    workbook = openpyxl.load_workbook(tmp_path / main.EXCEL_FILENAME)
    sheet = workbook[SHEET_PRICE_RESULTS]
    assert sheet.max_column == len(PRICE_RESULT_COLUMNS)
