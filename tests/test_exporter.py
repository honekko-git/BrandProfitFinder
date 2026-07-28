"""Phase 2 tests for excel.exporter."""

from decimal import Decimal
from pathlib import Path

import openpyxl

from excel.exporter import ExcelExporter, SHEET_PRICE_RESULTS
from excel.template import PRICE_RESULT_COLUMNS, PRODUCT_COLUMNS
from models.price_result import PriceResult
from models.product import Product
from price_compare.profit_calculator import ProfitCalculator


def test_export_products_creates_excel(tmp_path: Path) -> None:
    products = [
        Product(
            name="Sample Bag",
            brand="Demo Brand",
            price=500.0,
            currency="USD",
            exchange_rate=160.0,
            landed_cost=80000.0,
        )
    ]
    exporter = ExcelExporter(output_dir=tmp_path, filename="foundation.xlsx")
    output_path = exporter.export_products(products)

    assert output_path.exists()
    workbook = openpyxl.load_workbook(output_path)
    sheet = workbook.active
    assert sheet.cell(row=1, column=1).value == "name"
    assert sheet.cell(row=2, column=1).value == "Sample Bag"


def test_export_price_results_creates_sheet(tmp_path: Path) -> None:
    product = Product(name="Bag", brand="Brand", price=100.0, currency="USD", exchange_rate=100.0)
    result = ProfitCalculator().calculate(product, Decimal("80000"))
    exporter = ExcelExporter(output_dir=tmp_path, filename="phase2_results.xlsx")
    output_path = exporter.export_price_results([result])

    workbook = openpyxl.load_workbook(output_path)
    sheet = workbook.active
    assert sheet.title == SHEET_PRICE_RESULTS
    assert sheet.cell(row=1, column=1).value == "store_name"
    assert sheet.cell(row=2, column=3).value == "Bag"


def test_price_result_columns_exist(tmp_path: Path) -> None:
    product = Product(name="Bag", brand="Brand", price=100.0, currency="USD", exchange_rate=100.0)
    result = ProfitCalculator().calculate(product, Decimal("80000"))
    exporter = ExcelExporter(output_dir=tmp_path, filename="phase2_columns.xlsx")
    output_path = exporter.export_price_results([result])

    workbook = openpyxl.load_workbook(output_path)
    headers = [workbook.active.cell(row=1, column=i + 1).value for i in range(len(PRICE_RESULT_COLUMNS))]
    assert headers == PRICE_RESULT_COLUMNS


def test_numeric_columns_are_numeric(tmp_path: Path) -> None:
    product = Product(name="Bag", brand="Brand", price=100.0, currency="USD", exchange_rate=100.0)
    result = ProfitCalculator().calculate(product, Decimal("80000"))
    exporter = ExcelExporter(output_dir=tmp_path, filename="phase2_numeric.xlsx")
    output_path = exporter.export_price_results([result])

    workbook = openpyxl.load_workbook(output_path)
    sheet = workbook.active
    profit_col = PRICE_RESULT_COLUMNS.index("profit_jpy") + 1
    margin_col = PRICE_RESULT_COLUMNS.index("profit_margin") + 1
    roi_col = PRICE_RESULT_COLUMNS.index("roi") + 1
    assert isinstance(sheet.cell(row=2, column=profit_col).value, (int, float))
    assert isinstance(sheet.cell(row=2, column=margin_col).value, (int, float))
    assert isinstance(sheet.cell(row=2, column=roi_col).value, (int, float))


def test_url_hyperlink_and_filters(tmp_path: Path) -> None:
    product = Product(
        name="Bag",
        brand="Brand",
        price=100.0,
        currency="USD",
        exchange_rate=100.0,
        url="https://example.com/products/bag",
    )
    result = ProfitCalculator().calculate(product, Decimal("80000"))
    exporter = ExcelExporter(output_dir=tmp_path, filename="phase2_format.xlsx")
    output_path = exporter.export_price_results([result])

    workbook = openpyxl.load_workbook(output_path)
    sheet = workbook.active
    url_col = PRICE_RESULT_COLUMNS.index("url") + 1
    url_cell = sheet.cell(row=2, column=url_col)
    assert url_cell.hyperlink is not None
    assert sheet.auto_filter.ref is not None
    assert sheet.freeze_panes == "A2"


def test_profitability_and_error_message_output(tmp_path: Path) -> None:
    valid = ProfitCalculator().calculate(
        Product(name="Good", price=100.0, currency="USD", exchange_rate=100.0),
        Decimal("80000"),
    )
    invalid = ProfitCalculator().calculate(
        Product(name="Bad", price=0.0, currency="USD", exchange_rate=100.0),
        Decimal("80000"),
    )
    exporter = ExcelExporter(output_dir=tmp_path, filename="phase2_status.xlsx")
    output_path = exporter.export_price_results([valid, invalid])

    workbook = openpyxl.load_workbook(output_path)
    sheet = workbook.active
    profitable_col = PRICE_RESULT_COLUMNS.index("is_profitable") + 1
    status_col = PRICE_RESULT_COLUMNS.index("calculation_status") + 1
    error_col = PRICE_RESULT_COLUMNS.index("error_message") + 1
    assert sheet.cell(row=2, column=profitable_col).value is True
    assert sheet.cell(row=3, column=status_col).value != "success"
    assert sheet.cell(row=3, column=error_col).value


def test_export_phase2_workbook_has_product_and_result_sheets(tmp_path: Path) -> None:
    products = [Product(name="Bag", brand="Brand", price=100.0, currency="USD", exchange_rate=100.0)]
    results = [ProfitCalculator().calculate(products[0], Decimal("80000"))]
    exporter = ExcelExporter(output_dir=tmp_path, filename="phase2_workbook.xlsx")
    output_path = exporter.export_phase2_workbook(products, results)

    workbook = openpyxl.load_workbook(output_path)
    assert "Profit Ranking" in workbook.sheetnames
    assert SHEET_PRICE_RESULTS in workbook.sheetnames
    product_sheet = workbook["Profit Ranking"]
    assert product_sheet.cell(row=1, column=1).value == "name"
    assert len(PRODUCT_COLUMNS) == product_sheet.max_column
