"""Integration tests for GOAT marketplace."""

import json
from copy import deepcopy
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest

from config.constants import MARKETPLACE_GOAT, SHEET_DOMESTIC_LISTINGS
from excel.exporter import ExcelExporter
from excel.template import MARKETPLACE_LISTING_COLUMNS
from marketplace.goat_client import FakeGoatClient
from marketplace.goat_exceptions import GoatConfigurationError
from marketplace.goat_marketplace import GoatMarketplace, create_goat_marketplace
from marketplace.goat_settings import GoatSettings
from marketplace.marketplace_factory import create_marketplace
from models.product import Product
from price_compare.marketplace_profit_service import calculate_profit_from_search_results
from price_compare.profit_calculator import ProfitCalculator

FIXTURES = Path(__file__).parent / "fixtures"


def _settings(**overrides) -> GoatSettings:
    defaults = dict(
        enabled=True,
        timeout_seconds=10,
        page_size=20,
        max_pages=1,
        default_currency="JPY",
        demo_fixture_path="goat_search_normal.json",
        allow_unknown_currency=False,
        include_unavailable=False,
        include_sold=False,
        include_new=True,
        include_used=True,
        require_known_price=True,
    )
    defaults.update(overrides)
    return GoatSettings(**defaults)


def _marketplace(name: str = "goat_search_normal.json", **settings) -> GoatMarketplace:
    payload = json.loads((FIXTURES / name).read_text(encoding="utf-8"))
    return GoatMarketplace(
        client=FakeGoatClient(payload),
        settings=_settings(**settings),
    )


def test_marketplace_search() -> None:
    product = Product(
        name="Phase3 Sale Bag",
        brand="Gucci",
        model="GG-MARMONT",
        price=600.0,
        currency="USD",
        exchange_rate=150.0,
    )
    result = _marketplace().search(product)
    assert result.marketplace_name == MARKETPLACE_GOAT
    assert len(result.listings) == 3
    assert result.selected_price_jpy == Decimal("26500")


def test_factory_aliases() -> None:
    payload = json.loads((FIXTURES / "goat_search_normal.json").read_text(encoding="utf-8"))
    client = FakeGoatClient(payload)
    settings = _settings()
    for alias in ("goat_marketplace", "goat-marketplace"):
        marketplace = create_marketplace(alias, goat_client=client, goat_settings=settings)
        assert isinstance(marketplace, GoatMarketplace)


def test_factory_no_fake_without_client() -> None:
    with pytest.raises(GoatConfigurationError):
        create_marketplace("goat", goat_settings=_settings())


def test_profit_metadata() -> None:
    payload = json.loads((FIXTURES / "goat_search_normal.json").read_text(encoding="utf-8"))
    prada_only = {
        **payload,
        "items": [item for item in payload["items"] if item.get("brand") == "PRADA"],
        "total": 1,
    }
    product = Product(
        name="Phase3 Regular Shoes",
        brand="Prada",
        model="PRD-LOAFER",
        price=400.0,
        currency="EUR",
        exchange_rate=165.0,
    )
    marketplace = GoatMarketplace(client=FakeGoatClient(prada_only), settings=_settings())
    result = marketplace.search(product)
    price_result = calculate_profit_from_search_results([result], ProfitCalculator())[0]
    assert price_result.domestic_sale_price_jpy == Decimal("19800")
    assert price_result.metadata.get("fees_applied") is False
    assert price_result.metadata.get("source_style_code") == "P-DEMO-LOAFER-002"


def test_input_not_mutated() -> None:
    payload = json.loads((FIXTURES / "goat_search_normal.json").read_text(encoding="utf-8"))
    original = deepcopy(payload)
    marketplace = GoatMarketplace(client=FakeGoatClient(payload), settings=_settings())
    product = Product(name="Phase3 Sale Bag", brand="Gucci", model="GG-MARMONT", price=600.0)
    marketplace.search(product)
    assert payload == original


def test_deterministic_ordering() -> None:
    product = Product(name="Phase3 Sale Bag", brand="Gucci", model="GG-MARMONT", price=600.0)
    marketplace = _marketplace()
    first = marketplace.search(product)
    second = marketplace.search(product)
    assert [item.listing_id for item in first.listings] == [item.listing_id for item in second.listings]


def test_no_network_access() -> None:
    marketplace = _marketplace()
    product = Product(name="Phase3 Sale Bag", brand="Gucci", model="GG-MARMONT", price=600.0)
    with patch("utils.http.fetch_url") as mock_fetch, patch("utils.http.HttpClient") as mock_client:
        marketplace.search(product)
    mock_fetch.assert_not_called()
    mock_client.assert_not_called()


def test_excel_export_columns(tmp_path: Path) -> None:
    product = Product(name="Phase3 Sale Bag", brand="Gucci", model="GG-MARMONT", price=600.0)
    result = _marketplace().search(product)
    exporter = ExcelExporter(output_dir=tmp_path, filename="goat.xlsx")
    path = exporter.export_phase3_workbook([product], result.listings, [])
    workbook = openpyxl.load_workbook(path)
    sheet = workbook[SHEET_DOMESTIC_LISTINGS]
    headers = [sheet.cell(row=1, column=i + 1).value for i in range(sheet.max_column)]
    assert headers == MARKETPLACE_LISTING_COLUMNS
    assert "source_goat_style_code" in headers


def test_create_goat_marketplace_requires_client() -> None:
    with pytest.raises(GoatConfigurationError):
        create_goat_marketplace()
