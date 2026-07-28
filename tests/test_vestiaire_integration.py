"""Integration tests for Vestiaire marketplace."""

import json
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from config.constants import MARKETPLACE_VESTIAIRE, SHEET_DOMESTIC_LISTINGS
from excel.exporter import ExcelExporter
from excel.template import MARKETPLACE_LISTING_COLUMNS
from marketplace.marketplace_factory import create_marketplace
from marketplace.vestiaire_client import FakeVestiaireClient
from marketplace.vestiaire_exceptions import VestiaireConfigurationError
from marketplace.vestiaire_marketplace import VestiaireMarketplace, create_vestiaire_marketplace
from marketplace.vestiaire_settings import VestiaireSettings
from models.product import Product
from price_compare.marketplace_profit_service import calculate_profit_from_search_results
from price_compare.profit_calculator import ProfitCalculator

FIXTURES = Path(__file__).parent / "fixtures"


def _settings(**overrides) -> VestiaireSettings:
    defaults = dict(
        enabled=True,
        timeout_seconds=10,
        page_size=20,
        max_pages=2,
        default_currency="JPY",
        demo_fixture_path="vestiaire_search_normal.json",
        allow_unknown_currency=False,
        include_inactive=False,
        include_sold=False,
    )
    defaults.update(overrides)
    return VestiaireSettings(**defaults)


def _marketplace(name: str = "vestiaire_search_normal.json", **settings) -> VestiaireMarketplace:
    payload = json.loads((FIXTURES / name).read_text(encoding="utf-8"))
    return VestiaireMarketplace(
        client=FakeVestiaireClient(payload),
        settings=_settings(**settings),
    )


def test_marketplace_search() -> None:
    product = Product(name="Marmont Bag", brand="GUCCI", model="447632", price=600.0, currency="USD", exchange_rate=150.0)
    result = _marketplace().search(product)
    assert result.marketplace_name == MARKETPLACE_VESTIAIRE
    assert len(result.listings) == 1
    assert result.selected_price_jpy == Decimal("131500")


def test_empty_query() -> None:
    result = _marketplace().search(Product(name="", brand="", model=""), query="  ")
    assert result.status == "error"


def test_client_exception() -> None:
    mp = VestiaireMarketplace(client=FakeVestiaireClient(error=RuntimeError("fail")), settings=_settings())
    result = mp.search(Product(name="Bag", brand="GUCCI", model="447632", price=100.0))
    assert result.status == "error"


def test_empty_results() -> None:
    result = _marketplace("vestiaire_search_empty.json").search(Product(name="X", brand="Demo", price=100.0))
    assert result.status == "no_listings"


def test_pagination_two_pages() -> None:
    p1 = json.loads((FIXTURES / "vestiaire_search_page_1.json").read_text(encoding="utf-8"))
    p2 = json.loads((FIXTURES / "vestiaire_search_page_2.json").read_text(encoding="utf-8"))
    mp = VestiaireMarketplace(
        client=FakeVestiaireClient(pages={1: p1, 2: p2}),
        settings=_settings(max_pages=2),
    )
    result = mp.search(Product(name="Bag", brand="GUCCI", model="447632", price=100.0))
    assert len(result.listings) == 2


def test_duplicate_excluded_across_pages() -> None:
    p1 = json.loads((FIXTURES / "vestiaire_search_page_1.json").read_text(encoding="utf-8"))
    p2 = dict(p1)
    p2["page"] = 2
    mp = VestiaireMarketplace(
        client=FakeVestiaireClient(pages={1: p1, 2: p2}),
        settings=_settings(max_pages=2),
    )
    result = mp.search(Product(name="Bag", brand="GUCCI", model="447632", price=100.0))
    assert len(result.listings) == 1


def test_no_client_configured() -> None:
    result = VestiaireMarketplace(client=None, settings=_settings()).search(
        Product(name="Bag", brand="GUCCI", model="447632", price=100.0)
    )
    assert result.status == "error"
    assert "not configured" in (result.error_message or "")


def test_create_without_client_raises() -> None:
    with pytest.raises(VestiaireConfigurationError):
        create_vestiaire_marketplace(client=None)


def test_factory_with_client() -> None:
    payload = json.loads((FIXTURES / "vestiaire_search_normal.json").read_text(encoding="utf-8"))
    mp = create_marketplace(
        "vestiaire",
        vestiaire_client=FakeVestiaireClient(payload),
        vestiaire_settings=_settings(),
    )
    assert isinstance(mp, VestiaireMarketplace)


def test_factory_aliases() -> None:
    payload = json.loads((FIXTURES / "vestiaire_search_normal.json").read_text(encoding="utf-8"))
    client = FakeVestiaireClient(payload)
    settings = _settings()
    for alias in ("vestiaire_collective", "vestiaire-collective", "vc"):
        mp = create_marketplace(alias, vestiaire_client=client, vestiaire_settings=settings)
        assert isinstance(mp, VestiaireMarketplace)


def test_factory_no_fake_without_client() -> None:
    with pytest.raises(VestiaireConfigurationError):
        create_marketplace("vestiaire", vestiaire_settings=_settings())


def test_profit_not_auto_adjusted() -> None:
    product = Product(name="Bag", brand="GUCCI", model="447632", price=600.0, currency="USD", exchange_rate=150.0)
    result = _marketplace().search(product)
    price_result = calculate_profit_from_search_results([result], ProfitCalculator())[0]
    assert price_result.domestic_sale_price_jpy == Decimal("131500")
    assert price_result.metadata.get("adjustment_applied") is False


def test_excel_vestiaire_rows(tmp_path: Path) -> None:
    product = Product(name="Bag", brand="GUCCI", model="447632", price=600.0, currency="USD", exchange_rate=150.0)
    result = _marketplace("vestiaire_search_multiple.json").search(product)
    price_results = calculate_profit_from_search_results([result], ProfitCalculator())
    exporter = ExcelExporter(output_dir=tmp_path, filename="vestiaire.xlsx")
    output_path = exporter.export_phase3_workbook([product], result.listings, price_results)
    workbook = openpyxl.load_workbook(output_path)
    sheet = workbook[SHEET_DOMESTIC_LISTINGS]
    col = MARKETPLACE_LISTING_COLUMNS.index("marketplace_name") + 1
    assert MARKETPLACE_VESTIAIRE in [sheet.cell(row=r, column=col).value for r in range(2, sheet.max_row + 1)]
