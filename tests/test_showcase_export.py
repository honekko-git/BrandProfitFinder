"""Tests for showcase dashboard Excel export."""

from __future__ import annotations

from pathlib import Path

import openpyxl
from marketplace.yahoo_auction.client import FakeYahooAuctionClient
from profit_discovery.cli.discovery_command import (
    DiscoveryCommandOptions,
    export_discovery_report,
    run_discovery_command,
)
from profit_discovery.cli.demand_export import SHEET_DEMAND_OPPORTUNITY_RANKING
from profit_discovery.cli.opportunity_export import SHEET_OPPORTUNITY_RANKING
from profit_discovery.cli.showcase_export import SHOWCASE_COLUMNS, SHEET_SHOWCASE
from profit_discovery.discovery_runner import DiscoveryCandidateStatus, DiscoveryRunner
from profit_discovery.market_connector import SupplierMarketConnector
from profit_discovery.multi_brand import MultiBrandDiscoveryRunner
from profit_discovery.showcase import ShowcaseFormatter
from supplier.fashionphile.client import FashionphileClient


def _runner() -> MultiBrandDiscoveryRunner:
    supplier_client = FashionphileClient()
    discovery_runner = DiscoveryRunner(
        supplier_client=supplier_client,
        market_connector=SupplierMarketConnector(
            supplier_client=supplier_client,
            market_client=FakeYahooAuctionClient(),
        ),
    )
    return MultiBrandDiscoveryRunner(discovery_runner=discovery_runner)


def test_showcase_export_adds_showcase_sheet(tmp_path: Path) -> None:
    result = run_discovery_command(
        DiscoveryCommandOptions(
            brands=["Chanel"],
            keyword="wallet",
            max_results_per_brand=2,
        ),
        runner=_runner(),
    )

    export_path = export_discovery_report(result, output_dir=tmp_path)

    assert export_path is not None
    workbook = openpyxl.load_workbook(export_path)
    assert "Profit Analysis" in workbook.sheetnames
    assert SHEET_OPPORTUNITY_RANKING in workbook.sheetnames
    assert SHEET_DEMAND_OPPORTUNITY_RANKING in workbook.sheetnames
    assert SHEET_SHOWCASE in workbook.sheetnames

    sheet = workbook[SHEET_SHOWCASE]
    headers = [cell.value for cell in sheet[1]]
    assert headers == list(SHOWCASE_COLUMNS)
    showcase_count = len(ShowcaseFormatter().to_showcase_opportunities(result.ranked_demand_opportunities))
    assert sheet.max_row == showcase_count + 1


def test_showcase_export_preserves_existing_sheets(tmp_path: Path) -> None:
    result = run_discovery_command(
        DiscoveryCommandOptions(
            brands=["Chanel", "Louis Vuitton"],
            keyword="wallet",
            max_results_per_brand=1,
        ),
        runner=_runner(),
    )

    export_path = export_discovery_report(result, output_dir=tmp_path)

    assert export_path is not None
    workbook = openpyxl.load_workbook(export_path)
    profit_sheet = workbook["Profit Analysis"]
    assert profit_sheet.max_row >= 2

    successful = [
        candidate
        for candidate in result.ranked_candidates
        if candidate.status is DiscoveryCandidateStatus.SUCCESS and candidate.profit_result is not None
    ]
    assert successful
    assert profit_sheet.max_row == len(successful) + 1
