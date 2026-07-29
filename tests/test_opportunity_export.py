"""Tests for opportunity ranking Excel export."""

from __future__ import annotations

from pathlib import Path

import openpyxl
from marketplace.yahoo_auction.client import FakeYahooAuctionClient
from profit_discovery.cli.discovery_command import (
    DiscoveryCommandOptions,
    export_discovery_report,
    run_discovery_command,
)
from profit_discovery.cli.opportunity_export import (
    OPPORTUNITY_RANKING_COLUMNS,
    SHEET_OPPORTUNITY_RANKING,
)
from profit_discovery.discovery_runner import DiscoveryCandidateStatus, DiscoveryRunner
from profit_discovery.market_connector import SupplierMarketConnector
from profit_discovery.multi_brand import MultiBrandDiscoveryRunner
from supplier.fashionphile.client import FashionphileClient


def _discovery_runner() -> DiscoveryRunner:
    supplier_client = FashionphileClient()
    return DiscoveryRunner(
        supplier_client=supplier_client,
        market_connector=SupplierMarketConnector(
            supplier_client=supplier_client,
            market_client=FakeYahooAuctionClient(),
        ),
    )


def test_discovery_export_adds_opportunity_ranking_sheet(tmp_path: Path) -> None:
    result = run_discovery_command(
        DiscoveryCommandOptions(
            brands=["Chanel"],
            keyword="wallet",
            max_results_per_brand=2,
        ),
        runner=MultiBrandDiscoveryRunner(discovery_runner=_discovery_runner()),
    )

    export_path = export_discovery_report(result, output_dir=tmp_path)

    assert export_path is not None
    workbook = openpyxl.load_workbook(export_path)
    assert "Profit Analysis" in workbook.sheetnames
    assert SHEET_OPPORTUNITY_RANKING in workbook.sheetnames

    sheet = workbook[SHEET_OPPORTUNITY_RANKING]
    headers = [cell.value for cell in sheet[1]]
    assert headers == list(OPPORTUNITY_RANKING_COLUMNS)
    assert sheet.max_row == len(result.ranked_opportunities) + 1


def test_discovery_export_preserves_existing_profit_analysis_sheet(tmp_path: Path) -> None:
    result = run_discovery_command(
        DiscoveryCommandOptions(
            brands=["Chanel", "Louis Vuitton"],
            keyword="wallet",
            max_results_per_brand=1,
        ),
        runner=MultiBrandDiscoveryRunner(discovery_runner=_discovery_runner()),
    )

    export_path = export_discovery_report(result, output_dir=tmp_path)

    assert export_path is not None
    workbook = openpyxl.load_workbook(export_path)
    profit_sheet = workbook["Profit Analysis"]
    assert profit_sheet.max_row >= 2

    successful = [
        candidate
        for candidate in result.ranked_candidates
        if candidate.status is DiscoveryCandidateStatus.SUCCESS and candidate.profit_result is not None
    ]
    assert successful
    assert profit_sheet.max_row == len(successful) + 1
