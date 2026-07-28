"""Integration tests for StockX marketplace."""

import json
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from config.constants import MARKETPLACE_STOCKX, SHEET_DOMESTIC_LISTINGS
from excel.exporter import ExcelExporter
from excel.template import MARKETPLACE_LISTING_COLUMNS
from marketplace.marketplace_factory import create_marketplace
from marketplace.stockx_client import FakeStockXClient
from marketplace.stockx_exceptions import StockXConfigurationError
from marketplace.stockx_marketplace import StockXMarketplace, create_stockx_marketplace
from marketplace.stockx_settings import StockXSettings
from models.product import Product
from price_compare.marketplace_profit_service import calculate_profit_from_search_results
from price_compare.profit_calculator import ProfitCalculator

FIXTURES = Path(__file__).parent / "fixtures"


def _settings(**overrides) -> StockXSettings:
    defaults = dict(
        enabled=True,
        timeout_seconds=10,
        page_size=20,
        max_pages=2,
        default_currency="JPY",
        demo_fixture_path="stockx_search_normal.json",
        allow_unknown_currency=False,
        include_unavailable=False,
        include_sold=False,
        include_inactive=False,
        include_new=True,
        include_preowned=False,
        include_low_liquidity=True,
        require_known_lowest_ask=True,
        require_known_shipping=False,
        require_known_fees=False,
        minimum_sales_last_30_days=0,
        minimum_asks_count=0,
        minimum_bids_count=0,
        maximum_volatility_rate=None,
        preferred_price_source="LOWEST_ASK",
    )
    defaults.update(overrides)
    return StockXSettings(**defaults)


def _marketplace(name: str = "stockx_search_normal.json", **settings) -> StockXMarketplace:
    payload = json.loads((FIXTURES / name).read_text(encoding="utf-8"))
    return StockXMarketplace(
        client=FakeStockXClient(payload),
        settings=_settings(**settings),
    )


def test_marketplace_search() -> None:
    product = Product(
        name="Dunk Low",
        brand="NIKE",
        model="DD1391-100",
        price=120.0,
        currency="USD",
        exchange_rate=150.0,
    )
    result = _marketplace().search(product)
    assert result.marketplace_name == MARKETPLACE_STOCKX
    assert len(result.listings) == 3
    assert result.selected_price_jpy == Decimal("24500")


def test_factory_aliases() -> None:
    payload = json.loads((FIXTURES / "stockx_search_normal.json").read_text(encoding="utf-8"))
    client = FakeStockXClient(payload)
    settings = _settings()
    for alias in ("stock_x", "stock-x", "sx"):
        mp = create_marketplace(alias, stockx_client=client, stockx_settings=settings)
        assert isinstance(mp, StockXMarketplace)


def test_factory_no_fake_without_client() -> None:
    with pytest.raises(StockXConfigurationError):
        create_marketplace("stockx", stockx_settings=_settings())


def test_profit_metadata() -> None:
    product = Product(name="Dunk", brand="NIKE", model="DD1391-100", price=120.0, currency="USD", exchange_rate=150.0)
    result = _marketplace().search(product)
    price_result = calculate_profit_from_search_results([result], ProfitCalculator())[0]
    assert price_result.domestic_sale_price_jpy == Decimal("24500")
    assert price_result.metadata.get("market_price_applied") is False
    assert price_result.metadata.get("highest_bid_applied") is False
    assert price_result.metadata.get("last_sale_applied") is False
    assert price_result.metadata.get("fees_applied") is False


def test_excel_rows(tmp_path: Path) -> None:
    product = Product(name="Dunk", brand="NIKE", model="DD1391-100", price=120.0, currency="USD", exchange_rate=150.0)
    result = _marketplace().search(product)
    price_results = calculate_profit_from_search_results([result], ProfitCalculator())
    output_path = ExcelExporter(output_dir=tmp_path, filename="sx.xlsx").export_phase3_workbook(
        [product], result.listings, price_results
    )
    sheet = openpyxl.load_workbook(output_path)[SHEET_DOMESTIC_LISTINGS]
    col = MARKETPLACE_LISTING_COLUMNS.index("source_stockx_style_code") + 1
    assert "DD1391-100" in [sheet.cell(row=r, column=col).value for r in range(2, sheet.max_row + 1)]
