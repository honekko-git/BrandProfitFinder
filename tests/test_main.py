"""Foundation and Phase 2 tests for main entry point (no real network)."""

from pathlib import Path
from unittest.mock import patch

import openpyxl

from main import build_phase2_products, build_sample_products, resolve_marketplace_name, run


def test_build_sample_products_has_expected_values() -> None:
    products = build_sample_products()
    assert len(products) == 2
    assert all(product.landed_cost > 0 for product in products)
    assert all(product.profit != 0 or product.best_japanese_price() is None for product in products)


def test_build_phase2_products_has_local_data() -> None:
    products = build_phase2_products()
    assert len(products) == 3
    assert all(product.store_name for product in products)


def test_build_phase3_products_has_local_data() -> None:
    from main import build_phase3_products

    products = build_phase3_products()
    assert len(products) == 3


def test_resolve_marketplace_default_local(monkeypatch) -> None:
    monkeypatch.setattr("main.YAHOO_API_ENABLED", False)
    assert resolve_marketplace_name([]) == "local"


def test_resolve_marketplace_cli_yahoo() -> None:
    assert resolve_marketplace_name(["--marketplace", "yahoo"]) == "yahoo"


def test_resolve_marketplace_env_yahoo(monkeypatch) -> None:
    monkeypatch.setattr("main.YAHOO_API_ENABLED", True)
    assert resolve_marketplace_name([]) == "yahoo"


def test_run_exports_excel(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr("main.OUTPUT_DIR", tmp_path)
    monkeypatch.setattr("main.EXCEL_FILENAME", "foundation_run.xlsx")
    monkeypatch.setattr("main.YAHOO_API_ENABLED", False)

    output_path = run()
    assert output_path.exists()
    assert output_path.name == "foundation_run.xlsx"


def test_main_runs_without_network(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr("main.OUTPUT_DIR", tmp_path)
    monkeypatch.setattr("main.EXCEL_FILENAME", "phase2_main.xlsx")
    monkeypatch.setattr("main.YAHOO_API_ENABLED", False)

    with patch("utils.http.fetch_url") as mock_fetch, patch("utils.http.HttpClient") as mock_client:
        output_path = run()

    assert output_path.exists()
    mock_fetch.assert_not_called()
    mock_client.assert_not_called()
    workbook = openpyxl.load_workbook(output_path)
    assert "Profit Analysis" in workbook.sheetnames
    assert "Domestic Listings" in workbook.sheetnames


def test_yahoo_enabled_without_client_id_falls_back(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr("main.OUTPUT_DIR", tmp_path)
    monkeypatch.setattr("main.EXCEL_FILENAME", "yahoo_fallback.xlsx")
    monkeypatch.setattr("main.YAHOO_API_ENABLED", True)
    monkeypatch.setattr("config.settings.YAHOO_CLIENT_ID", "")

    with patch("utils.http.fetch_url") as mock_fetch, patch("utils.http.HttpClient"):
        output_path = run()

    assert output_path.exists()
    mock_fetch.assert_not_called()


def test_main_runnable_twice(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr("main.OUTPUT_DIR", tmp_path)
    monkeypatch.setattr("main.EXCEL_FILENAME", "phase3_twice.xlsx")
    monkeypatch.setattr("main.YAHOO_API_ENABLED", False)

    with patch("utils.http.fetch_url"), patch("utils.http.HttpClient"):
        first = run()
        second = run()

    assert first.exists()
    assert second.exists()
