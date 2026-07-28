"""Integration tests for Farfetch marketplace."""

import json
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from config.constants import MARKETPLACE_FARFETCH, SHEET_DOMESTIC_LISTINGS
from excel.exporter import ExcelExporter
from excel.template import MARKETPLACE_LISTING_COLUMNS
from marketplace.farfetch_client import FakeFarfetchClient
from marketplace.farfetch_exceptions import FarfetchConfigurationError
from marketplace.farfetch_marketplace import FarfetchMarketplace, create_farfetch_marketplace
from marketplace.farfetch_settings import FarfetchSettings
from marketplace.marketplace_factory import create_marketplace
from models.product import Product
from price_compare.marketplace_profit_service import calculate_profit_from_search_results
from price_compare.profit_calculator import ProfitCalculator

FIXTURES = Path(__file__).parent / "fixtures"


def _settings(**overrides) -> FarfetchSettings:
    defaults = dict(
        enabled=True,
        timeout_seconds=10,
        page_size=20,
        max_pages=2,
        default_currency="JPY",
        demo_fixture_path="farfetch_search_normal.json",
        allow_unknown_currency=False,
        include_unavailable=False,
        include_sold=False,
        include_reserved=False,
        include_discounted_only=False,
        include_full_price_only=False,
        include_low_stock=True,
        include_final_sale=False,
        include_partner_boutiques=True,
        include_platform_inventory=True,
        require_known_shipping=False,
        require_known_duties=False,
    )
    defaults.update(overrides)
    return FarfetchSettings(**defaults)


def _marketplace(name: str = "farfetch_search_normal.json", **settings) -> FarfetchMarketplace:
    payload = json.loads((FIXTURES / name).read_text(encoding="utf-8"))
    return FarfetchMarketplace(
        client=FakeFarfetchClient(payload),
        settings=_settings(**settings),
    )


def test_marketplace_search() -> None:
    product = Product(
        name="GG Marmont",
        brand="GUCCI",
        model="447632",
        price=12000.0,
        currency="USD",
        exchange_rate=150.0,
    )
    result = _marketplace().search(product)
    assert result.marketplace_name == MARKETPLACE_FARFETCH
    assert len(result.listings) == 3
    assert result.selected_price_jpy == Decimal("298000")


def test_factory_aliases() -> None:
    payload = json.loads((FIXTURES / "farfetch_search_normal.json").read_text(encoding="utf-8"))
    client = FakeFarfetchClient(payload)
    settings = _settings()
    for alias in ("far_fetch", "far-fetch", "ff"):
        mp = create_marketplace(alias, farfetch_client=client, farfetch_settings=settings)
        assert isinstance(mp, FarfetchMarketplace)


def test_factory_no_fake_without_client() -> None:
    with pytest.raises(FarfetchConfigurationError):
        create_marketplace("farfetch", farfetch_settings=_settings())


def test_profit_metadata() -> None:
    product = Product(name="Bag", brand="GUCCI", model="447632", price=12000.0, currency="USD", exchange_rate=150.0)
    result = _marketplace().search(product)
    price_result = calculate_profit_from_search_results([result], ProfitCalculator())[0]
    assert price_result.domestic_sale_price_jpy == Decimal("298000")
    assert price_result.metadata.get("discount_applied") is False
    assert price_result.metadata.get("duties_applied") is False
    assert price_result.metadata.get("variant_price_applied") is False
    assert price_result.metadata.get("source_style_code") == "447632-DTD1T-1000"


def test_excel_rows(tmp_path: Path) -> None:
    product = Product(name="Bag", brand="GUCCI", model="447632", price=12000.0, currency="USD", exchange_rate=150.0)
    result = _marketplace().search(product)
    price_results = calculate_profit_from_search_results([result], ProfitCalculator())
    output_path = ExcelExporter(output_dir=tmp_path, filename="ff.xlsx").export_phase3_workbook(
        [product], result.listings, price_results
    )
    sheet = openpyxl.load_workbook(output_path)[SHEET_DOMESTIC_LISTINGS]
    col = MARKETPLACE_LISTING_COLUMNS.index("source_style_code") + 1
    assert "447632-DTD1T-1000" in [sheet.cell(row=r, column=col).value for r in range(2, sheet.max_row + 1)]


def test_pagination() -> None:
    p1 = json.loads((FIXTURES / "farfetch_search_page_1.json").read_text(encoding="utf-8"))
    p2 = json.loads((FIXTURES / "farfetch_search_page_2.json").read_text(encoding="utf-8"))
    mp = FarfetchMarketplace(
        client=FakeFarfetchClient(pages={1: p1, 2: p2}),
        settings=_settings(max_pages=2),
    )
    product = Product(name="Bag", brand="GUCCI", model="447632", price=100.0, currency="USD", exchange_rate=150.0)
    result = mp.search(product)
    assert len(result.listings) == 3
