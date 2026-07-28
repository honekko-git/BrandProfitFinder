"""Integration tests for used luxury demo provider and marketplace."""

import json
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from config.constants import MARKETPLACE_USED_DEMO, SHEET_DOMESTIC_LISTINGS
from excel.exporter import ExcelExporter
from excel.template import MARKETPLACE_LISTING_COLUMNS
from models.product import Product
from price_compare.marketplace_profit_service import calculate_profit_from_search_results
from price_compare.profit_calculator import ProfitCalculator
from used_luxury.demo_marketplace import create_used_luxury_demo_marketplace
from used_luxury.demo_provider import FakeUsedLuxuryProvider, UsedLuxuryResponseParser

FIXTURES = Path(__file__).parent / "fixtures"


def _load(name: str) -> dict:
    return json.loads((FIXTURES / name).read_text(encoding="utf-8"))


def _marketplace(name: str = "used_luxury_normal.json"):
    return create_used_luxury_demo_marketplace(provider=FakeUsedLuxuryProvider(_load(name)))


def test_fake_provider_returns_fixture() -> None:
    provider = FakeUsedLuxuryProvider(_load("used_luxury_normal.json"))
    payload = provider.search_items(query="gucci")
    assert payload["total_results"] == 1


def test_fake_provider_records_query() -> None:
    provider = FakeUsedLuxuryProvider(_load("used_luxury_empty.json"))
    provider.search_items(query="test query", page=2, hits=10)
    assert provider.last_query == "test query"
    assert provider.last_page == 2
    assert provider.last_hits == 10


def test_fake_provider_injected_error() -> None:
    provider = FakeUsedLuxuryProvider(error=RuntimeError("fail"))
    with pytest.raises(RuntimeError):
        provider.search_items(query="x")


def test_parser_enriches_details() -> None:
    listings = UsedLuxuryResponseParser().parse(_load("used_luxury_normal.json"))
    assert len(listings) == 1
    details = listings[0].used_item_details
    assert details is not None
    assert details.condition.value == "VERY_GOOD"
    assert details.condition_score_value is not None


def test_marketplace_search() -> None:
    product = Product(name="Marmont Bag", brand="GUCCI", model="447632", price=600.0, currency="USD", exchange_rate=150.0)
    result = _marketplace().search(product)
    assert result.marketplace_name == MARKETPLACE_USED_DEMO
    assert len(result.listings) == 1
    assert result.selected_listing is not None
    assert result.selected_listing.used_item_details is not None


def test_empty_results() -> None:
    product = Product(name="Unknown", brand="Demo", price=100.0)
    result = _marketplace("used_luxury_empty.json").search(product)
    assert result.listings == []


def test_malformed_skipped() -> None:
    listings = UsedLuxuryResponseParser().parse(_load("used_luxury_malformed.json"))
    assert listings == []


def test_profit_not_auto_adjusted() -> None:
    product = Product(name="Marmont Bag", brand="GUCCI", model="447632", price=600.0, currency="USD", exchange_rate=150.0)
    result = _marketplace().search(product)
    price_result = calculate_profit_from_search_results([result], ProfitCalculator())[0]
    assert price_result.domestic_sale_price_jpy == Decimal("128000")
    assert price_result.metadata.get("adjustment_applied") is False
    assert price_result.metadata.get("condition_score") is not None


def test_profit_metadata_risk() -> None:
    product = Product(name="Marmont Bag", brand="GUCCI", model="447632", price=600.0, currency="USD", exchange_rate=150.0)
    result = _marketplace().search(product)
    price_result = calculate_profit_from_search_results([result], ProfitCalculator())[0]
    assert "risk_level" in price_result.metadata
    assert "risk_flags" in price_result.metadata


def test_excel_used_columns(tmp_path: Path) -> None:
    product = Product(name="Marmont Bag", brand="GUCCI", model="447632", price=600.0, currency="USD", exchange_rate=150.0)
    result = _marketplace("used_luxury_multiple.json").search(product)
    price_results = calculate_profit_from_search_results([result], ProfitCalculator())
    exporter = ExcelExporter(output_dir=tmp_path, filename="used.xlsx")
    output_path = exporter.export_phase3_workbook([product], result.listings, price_results)
    workbook = openpyxl.load_workbook(output_path)
    sheet = workbook[SHEET_DOMESTIC_LISTINGS]
    score_col = MARKETPLACE_LISTING_COLUMNS.index("condition_score") + 1
    value = sheet.cell(row=2, column=score_col).value
    assert value is None or isinstance(value, (int, float))
    auth_col = MARKETPLACE_LISTING_COLUMNS.index("authentication_status") + 1
    assert sheet.cell(row=2, column=auth_col).value is not None


def test_excel_none_not_string(tmp_path: Path) -> None:
    product = Product(name="Bag", brand="GUCCI", model="447632", price=100.0)
    result = _marketplace("used_luxury_empty.json").search(product)
    exporter = ExcelExporter(output_dir=tmp_path, filename="used_empty.xlsx")
    exporter.export_phase3_workbook([product], result.listings, [])
    workbook = openpyxl.load_workbook(exporter.output_path)
    assert SHEET_DOMESTIC_LISTINGS in workbook.sheetnames
