"""End-to-end tests for discovery CLI opportunity pipeline."""

from __future__ import annotations

from pathlib import Path

import openpyxl
from marketplace.yahoo_auction.client import FakeYahooAuctionClient
from profit_discovery.cli.discovery_command import (
    DiscoveryCommandOptions,
    export_discovery_report,
    run_discovery_command,
)
from profit_discovery.cli.opportunity_export import SHEET_OPPORTUNITY_RANKING
from profit_discovery.discovery_runner import DiscoveryCandidateStatus, DiscoveryRunner
from profit_discovery.market_connector import SupplierMarketConnector
from profit_discovery.multi_brand import MultiBrandDiscoveryRequest, MultiBrandDiscoveryRunner
from supplier.fashionphile.client import FashionphileClient


def _discovery_runner() -> DiscoveryRunner:
    supplier_client = FashionphileClient()
    return DiscoveryRunner(
        supplier_client=supplier_client,
        market_connector=SupplierMarketConnector(
            supplier_client=supplier_client,
            market_client=FakeYahooAuctionClient(),
        ),
    )


def test_full_cli_discovery_opportunity_export_pipeline(tmp_path: Path) -> None:
    discovery_runner = _discovery_runner()
    runner = MultiBrandDiscoveryRunner(discovery_runner=discovery_runner)
    options = DiscoveryCommandOptions(
        brands=["Chanel"],
        keyword="wallet",
        max_results_per_brand=2,
        export=True,
    )

    direct_result = runner.run(
        MultiBrandDiscoveryRequest(
            brands=options.brands,
            keyword=options.keyword,
            max_results_per_brand=options.max_results_per_brand,
        ),
    )
    cli_result = run_discovery_command(
        options,
        runner=runner,
    )
    export_path = export_discovery_report(cli_result, output_dir=tmp_path)

    direct_candidate = next(
        candidate
        for candidate in direct_result.ranked_candidates
        if candidate.status is DiscoveryCandidateStatus.SUCCESS
    )
    cli_candidate = next(
        item.candidate
        for item in cli_result.ranked_opportunities
        if item.candidate.status is DiscoveryCandidateStatus.SUCCESS
    )

    assert direct_candidate.profit_result is not None
    assert cli_candidate.profit_result is not None
    assert cli_candidate.profit_result.profit_jpy == direct_candidate.profit_result.profit_jpy
    assert cli_candidate.profit_result.roi == direct_candidate.profit_result.roi
    assert cli_candidate.buy_decision is not None
    assert direct_candidate.buy_decision is not None
    assert cli_candidate.buy_decision.decision == direct_candidate.buy_decision.decision

    assert cli_result.ranked_opportunities
    assert len(cli_result.ranked_opportunities) == len(direct_result.ranked_opportunities)
    assert [
        item.candidate.supplier_product.external_id for item in cli_result.ranked_opportunities
    ] == [
        item.candidate.supplier_product.external_id for item in direct_result.ranked_opportunities
    ]
    assert export_path is not None
    workbook = openpyxl.load_workbook(export_path)
    assert "Profit Analysis" in workbook.sheetnames
    assert SHEET_OPPORTUNITY_RANKING in workbook.sheetnames

    assert cli_result.total_candidates == direct_result.total_candidates
