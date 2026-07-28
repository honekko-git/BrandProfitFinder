"""Integration tests for Yahoo marketplace factory and profit pipeline."""

import json
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

import httpx
import openpyxl
import pytest

from config.constants import SHEET_DOMESTIC_LISTINGS
from excel.exporter import ExcelExporter, SHEET_PRICE_RESULTS
from excel.template import PRODUCT_COLUMNS
from marketplace.base_marketplace import BaseMarketplace
from marketplace.marketplace_factory import create_marketplace, get_all_marketplaces
from marketplace.yahoo_api_client import YahooApiClient
from marketplace.yahoo_marketplace import YahooMarketplace
from marketplace.yahoo_settings import YahooApiSettings
from models.product import Product
from price_compare.marketplace_profit_service import calculate_profit_from_search_results
from price_compare.price_comparator import PriceComparator, PriceSelectionStrategy
from price_compare.profit_calculator import ProfitCalculator
from price_compare.ranking_engine import RankingEngine, RankingSortKey

FIXTURES = Path(__file__).parent / "fixtures"


def _yahoo_settings(**overrides) -> YahooApiSettings:
    defaults = dict(
        client_id="dummy-test-client-id",
        base_url="https://shopping.yahooapis.jp/ShoppingWebService/V3/itemSearch",
        timeout_seconds=10,
        results=20,
        enabled=True,
    )
    defaults.update(overrides)
    return YahooApiSettings(**defaults)


def _yahoo_marketplace() -> YahooMarketplace:
    payload = json.loads((FIXTURES / "yahoo_item_search_success.json").read_text(encoding="utf-8"))

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, json=payload)

    client = YahooApiClient(settings=_yahoo_settings(), client=httpx.Client(transport=httpx.MockTransport(handler)))
    return YahooMarketplace(client=client)


def test_create_yahoo_marketplace() -> None:
    marketplace = create_marketplace("yahoo", yahoo_settings=_yahoo_settings())
    assert isinstance(marketplace, YahooMarketplace)
    assert isinstance(marketplace, BaseMarketplace)


def test_yahoo_case_insensitive() -> None:
    marketplace = create_marketplace("YAHOO", yahoo_settings=_yahoo_settings())
    assert isinstance(marketplace, YahooMarketplace)


def test_yahoo_whitespace() -> None:
    marketplace = create_marketplace("  yahoo  ", yahoo_settings=_yahoo_settings())
    assert isinstance(marketplace, YahooMarketplace)


def test_local_still_works() -> None:
    marketplace = create_marketplace("local")
    assert marketplace.marketplace_name == "local"


def test_unsupported_marketplace() -> None:
    with pytest.raises(ValueError, match="Unsupported marketplace"):
        create_marketplace("ebay")


def test_get_all_marketplaces_includes_yahoo() -> None:
    marketplaces = get_all_marketplaces(yahoo_settings=_yahoo_settings())
    names = [item.marketplace_name for item in marketplaces]
    assert "local" in names
    assert "yahoo" in names
    assert "amazon_jp" in names
    assert "rakuten" in names
    assert len(marketplaces) == 5


def test_factory_import_without_api_key() -> None:
    settings = _yahoo_settings(client_id="")
    marketplace = create_marketplace("yahoo", yahoo_settings=settings)
    assert isinstance(marketplace, YahooMarketplace)


def test_prices_from_listings() -> None:
    marketplace = _yahoo_marketplace()
    product = Product(name="Gucci Marmont Bag", brand="Gucci", model="GG-MARMONT", sku="4901234567890", price=600.0)
    result = marketplace.search(product)
    comparator = PriceComparator()
    prices = comparator.prices_from_listings(result.valid_listings)
    assert len(prices) >= 2


def test_price_comparator_lowest() -> None:
    marketplace = _yahoo_marketplace()
    product = Product(name="Gucci Marmont Bag", brand="Gucci", model="GG-MARMONT", sku="4901234567890", price=600.0)
    result = marketplace.search(product)
    selected = PriceComparator().select_from_listings(result.valid_listings, strategy=PriceSelectionStrategy.LOWEST)
    assert selected is not None
    assert selected[1] == Decimal("88000")


def test_price_comparator_highest() -> None:
    marketplace = _yahoo_marketplace()
    product = Product(name="Gucci Marmont Bag", brand="Gucci", model="GG-MARMONT", sku="4901234567890", price=600.0)
    result = marketplace.search(product)
    assert result.selected_price_jpy == Decimal("102000")


def test_marketplace_profit_service() -> None:
    marketplace = _yahoo_marketplace()
    product = Product(name="Gucci Marmont Bag", brand="Gucci", model="GG-MARMONT", sku="4901234567890", price=600.0, currency="USD", exchange_rate=150.0)
    result = marketplace.search(product)
    calculator = ProfitCalculator()
    price_results = calculate_profit_from_search_results([result], calculator)
    assert len(price_results) == 1
    assert price_results[0].domestic_sale_price_jpy == Decimal("102000")


def test_ranking_engine_integration() -> None:
    marketplace = _yahoo_marketplace()
    product = Product(name="Gucci Marmont Bag", brand="Gucci", model="GG-MARMONT", sku="4901234567890", price=600.0, currency="USD", exchange_rate=150.0)
    result = marketplace.search(product)
    calculator = ProfitCalculator()
    price_result = calculate_profit_from_search_results([result], calculator)[0]
    ranked = RankingEngine(calculator.config).rank([price_result], sort_key=RankingSortKey.PROFIT, descending=True)
    assert len(ranked) == 1


def test_excel_domestic_listings_yahoo(tmp_path: Path) -> None:
    marketplace = _yahoo_marketplace()
    product = Product(name="Gucci Marmont Bag", brand="Gucci", model="GG-MARMONT", sku="4901234567890", price=600.0, currency="USD", exchange_rate=150.0)
    result = marketplace.search(product)
    calculator = ProfitCalculator()
    price_results = calculate_profit_from_search_results([result], calculator)
    exporter = ExcelExporter(output_dir=tmp_path, filename="yahoo.xlsx")
    output_path = exporter.export_phase3_workbook([product], result.listings, price_results)

    workbook = openpyxl.load_workbook(output_path)
    sheet = workbook[SHEET_DOMESTIC_LISTINGS]
    assert sheet.cell(row=2, column=1).value == "yahoo"
    url_cell = sheet.cell(row=2, column=14)
    assert url_cell.hyperlink is not None


def test_excel_jan_leading_zero(tmp_path: Path) -> None:
    marketplace = _yahoo_marketplace()
    product = Product(name="Gucci Marmont Bag", brand="Gucci", model="GG-MARMONT", sku="4901234567890", price=600.0)
    result = marketplace.search(product)
    exporter = ExcelExporter(output_dir=tmp_path, filename="yahoo_jan.xlsx")
    exporter.export_phase3_workbook([], result.listings, [])
    workbook = openpyxl.load_workbook(exporter.output_path)
    sheet = workbook[SHEET_DOMESTIC_LISTINGS]
    jan_col = 7
    values = [sheet.cell(row=row, column=jan_col).value for row in range(2, sheet.max_row + 1)]
    assert "0123456789012" in values


def test_excel_zero_listings(tmp_path: Path) -> None:
    payload = json.loads((FIXTURES / "yahoo_item_search_empty.json").read_text(encoding="utf-8"))

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, json=payload)

    client = YahooApiClient(settings=_yahoo_settings(), client=httpx.Client(transport=httpx.MockTransport(handler)))
    marketplace = YahooMarketplace(client=client)
    product = Product(name="Unknown", brand="Demo", price=100.0)
    result = marketplace.search(product)
    exporter = ExcelExporter(output_dir=tmp_path, filename="yahoo_empty.xlsx")
    output_path = exporter.export_phase3_workbook([product], result.listings, [])
    workbook = openpyxl.load_workbook(output_path)
    assert SHEET_DOMESTIC_LISTINGS in workbook.sheetnames


def test_existing_product_sheet_preserved(tmp_path: Path) -> None:
    product = Product(name="Sample Bag", brand="Demo", price=500.0)
    exporter = ExcelExporter(output_dir=tmp_path, filename="legacy.xlsx")
    output_path = exporter.export_products([product])
    workbook = openpyxl.load_workbook(output_path)
    assert workbook.active.cell(row=1, column=1).value == "name"
    assert len(PRODUCT_COLUMNS) == workbook.active.max_column


def test_existing_price_result_sheet_preserved(tmp_path: Path) -> None:
    marketplace = _yahoo_marketplace()
    product = Product(name="Gucci Marmont Bag", brand="Gucci", model="GG-MARMONT", sku="4901234567890", price=600.0, currency="USD", exchange_rate=150.0)
    result = marketplace.search(product)
    calculator = ProfitCalculator()
    price_results = calculate_profit_from_search_results([result], calculator)
    exporter = ExcelExporter(output_dir=tmp_path, filename="yahoo_price.xlsx")
    exporter.export_phase3_workbook([product], result.listings, price_results)
    workbook = openpyxl.load_workbook(exporter.output_path)
    assert SHEET_PRICE_RESULTS in workbook.sheetnames
