"""Tests for market source truth in showcase Excel export."""

from __future__ import annotations

import io
from pathlib import Path

import openpyxl

from marketplace.domestic_market import FakeYahooAuctionTransport
from profit_discovery.cli.discovery_command import (
    DiscoveryCommandOptions,
    build_default_discovery_runner,
    build_domestic_market_runtime_config,
    export_discovery_report,
    read_runtime_market_execution,
    run_discovery_command,
)
from profit_discovery.cli.showcase_export import SHOWCASE_COLUMNS, SHEET_SHOWCASE


def _live_runner():
    transport = FakeYahooAuctionTransport(
        default_items=[
            {"title": "Chanel Classic Wallet Black Caviar", "sold_price": 155000},
            {"title": "Chanel CC Wallet Medium", "sold_price": 160000},
        ]
    )
    return build_default_discovery_runner(
        market_config=build_domestic_market_runtime_config(live_market=True),
        injected_transport=transport,
    ), transport


def test_showcase_export_includes_market_truth_columns(tmp_path: Path) -> None:
    runner, transport = _live_runner()
    result = run_discovery_command(
        DiscoveryCommandOptions(
            brands=["CHANEL"],
            keyword="wallet",
            max_results_per_brand=2,
            live_market=True,
        ),
        runner=runner,
        injected_transport=transport,
        output=io.StringIO(),
    )

    export_path = export_discovery_report(
        result,
        output_dir=tmp_path,
        market_execution=read_runtime_market_execution(runner),
    )
    assert export_path is not None

    workbook = openpyxl.load_workbook(export_path)
    assert SHEET_SHOWCASE in workbook.sheetnames

    sheet = workbook[SHEET_SHOWCASE]
    headers = [cell.value for cell in sheet[1]]
    assert headers == list(SHOWCASE_COLUMNS)
    assert "Requested Mode" in headers
    assert "Actual Source" in headers
    assert "Fallback" in headers

    if sheet.max_row >= 2:
        row = {headers[index]: sheet.cell(row=2, column=index + 1).value for index in range(len(headers))}
        assert row["Requested Mode"] == "LIVE"
        assert row["Actual Source"] == "Yahoo Auction LIVE"
        assert row["Fallback"] == "NO"
