"""Unit tests for domestic marketplace Excel export."""

from decimal import Decimal
from pathlib import Path

import openpyxl

from config.constants import SHEET_DOMESTIC_LISTINGS
from excel.exporter import ExcelExporter, SHEET_PRICE_RESULTS
from excel.template import MARKETPLACE_LISTING_COLUMNS, PRICE_RESULT_COLUMNS, PRODUCT_COLUMNS
from models.marketplace_listing import MarketplaceListing
from models.product import Product
from price_compare.profit_calculator import ProfitCalculator


def test_export_domestic_listings_sheet(tmp_path: Path) -> None:
    listings = [
        MarketplaceListing(
            marketplace_name="local-yahoo",
            listing_id="Y-1",
            title="Bag",
            price_jpy=Decimal("10000"),
            shipping_jpy=Decimal("500"),
            listing_url="https://example.com/item",
        )
    ]
    products = [Product(name="Bag", brand="Gucci", price=100.0, currency="USD", exchange_rate=100.0)]
    results = [ProfitCalculator().calculate(products[0], Decimal("10000"))]
    exporter = ExcelExporter(output_dir=tmp_path, filename="phase3.xlsx")
    output_path = exporter.export_phase3_workbook(products, listings, results)

    workbook = openpyxl.load_workbook(output_path)
    assert SHEET_DOMESTIC_LISTINGS in workbook.sheetnames
    assert SHEET_PRICE_RESULTS in workbook.sheetnames
    assert "Profit Ranking" in workbook.sheetnames


def test_listing_columns_exist(tmp_path: Path) -> None:
    exporter = ExcelExporter(output_dir=tmp_path, filename="phase3_columns.xlsx")
    output_path = exporter.export_phase3_workbook([], [], [])
    workbook = openpyxl.load_workbook(output_path)
    sheet = workbook[SHEET_DOMESTIC_LISTINGS]
    headers = [sheet.cell(row=1, column=i + 1).value for i in range(len(MARKETPLACE_LISTING_COLUMNS))]
    assert headers == MARKETPLACE_LISTING_COLUMNS


def test_listing_numeric_columns(tmp_path: Path) -> None:
    listings = [
        MarketplaceListing(
            marketplace_name="local",
            title="Item",
            price_jpy=Decimal("10000"),
            shipping_jpy=Decimal("500"),
            listing_url="https://example.com/item",
        )
    ]
    exporter = ExcelExporter(output_dir=tmp_path, filename="phase3_numeric.xlsx")
    exporter.export_phase3_workbook([], listings, [])
    workbook = openpyxl.load_workbook(exporter.output_path)
    sheet = workbook[SHEET_DOMESTIC_LISTINGS]
    price_col = MARKETPLACE_LISTING_COLUMNS.index("price_jpy") + 1
    assert isinstance(sheet.cell(row=2, column=price_col).value, (int, float))


def test_listing_url_hyperlink_and_filters(tmp_path: Path) -> None:
    listings = [
        MarketplaceListing(
            marketplace_name="local",
            title="Item",
            price_jpy=Decimal("10000"),
            listing_url="https://example.com/item",
        )
    ]
    exporter = ExcelExporter(output_dir=tmp_path, filename="phase3_links.xlsx")
    exporter.export_phase3_workbook([], listings, [])
    workbook = openpyxl.load_workbook(exporter.output_path)
    sheet = workbook[SHEET_DOMESTIC_LISTINGS]
    url_col = MARKETPLACE_LISTING_COLUMNS.index("listing_url") + 1
    cell = sheet.cell(row=2, column=url_col)
    assert cell.hyperlink is not None
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref is not None


def test_zero_listings_export(tmp_path: Path) -> None:
    exporter = ExcelExporter(output_dir=tmp_path, filename="phase3_empty.xlsx")
    output_path = exporter.export_phase3_workbook([], [], [])
    workbook = openpyxl.load_workbook(output_path)
    sheet = workbook[SHEET_DOMESTIC_LISTINGS]
    assert sheet.cell(row=1, column=1).value == "marketplace_name"
    assert sheet.max_row == 1


def test_existing_product_export_preserved(tmp_path: Path) -> None:
    products = [Product(name="Sample Bag", brand="Demo", price=500.0)]
    exporter = ExcelExporter(output_dir=tmp_path, filename="legacy.xlsx")
    output_path = exporter.export_products(products)
    workbook = openpyxl.load_workbook(output_path)
    sheet = workbook.active
    assert sheet.cell(row=1, column=1).value == "name"
    assert len(PRODUCT_COLUMNS) == sheet.max_column


def test_existing_price_result_export_preserved(tmp_path: Path) -> None:
    product = Product(name="Bag", price=100.0, currency="USD", exchange_rate=100.0)
    result = ProfitCalculator().calculate(product, Decimal("80000"))
    exporter = ExcelExporter(output_dir=tmp_path, filename="price_result.xlsx")
    exporter.export_price_results([result])
    workbook = openpyxl.load_workbook(exporter.output_path)
    headers = [workbook.active.cell(row=1, column=i + 1).value for i in range(len(PRICE_RESULT_COLUMNS))]
    assert headers == PRICE_RESULT_COLUMNS
