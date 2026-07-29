"""Consistency tests across CLI, runner result, and export ordering."""

from __future__ import annotations

from pathlib import Path

import openpyxl
from marketplace.yahoo_auction.client import FakeYahooAuctionClient
from profit_discovery.cli.discovery_command import (
    DiscoveryCommandOptions,
    export_discovery_report,
    run_discovery_command,
)
from profit_discovery.cli.opportunity_export import SHEET_OPPORTUNITY_RANKING
from profit_discovery.discovery_runner import DiscoveryCandidateStatus, DiscoveryRunner
from profit_discovery.market_connector import SupplierMarketConnector
from profit_discovery.multi_brand import MultiBrandDiscoveryRunner
from supplier.fashionphile.client import FashionphileClient


def _runner() -> MultiBrandDiscoveryRunner:
    supplier_client = FashionphileClient()
    discovery_runner = DiscoveryRunner(
        supplier_client=supplier_client,
        market_connector=SupplierMarketConnector(
            supplier_client=supplier_client,
            market_client=FakeYahooAuctionClient(),
        ),
    )
    return MultiBrandDiscoveryRunner(discovery_runner=discovery_runner)


def test_cli_runner_and_export_share_opportunity_order(tmp_path: Path) -> None:
    runner = _runner()
    cli_result = run_discovery_command(
        DiscoveryCommandOptions(
            brands=["Chanel", "Louis Vuitton"],
            keyword="wallet",
            max_results_per_brand=2,
        ),
        runner=runner,
    )

    runner_order = [
        item.candidate.supplier_product.external_id for item in cli_result.ranked_opportunities
    ]
    export_path = export_discovery_report(cli_result, output_dir=tmp_path)
    assert export_path is not None

    workbook = openpyxl.load_workbook(export_path)
    sheet = workbook[SHEET_OPPORTUNITY_RANKING]
    export_order = [
        sheet.cell(row=row_index, column=2).value
        for row_index in range(2, sheet.max_row + 1)
    ]

    assert export_order == [
        item.candidate.supplier_product.title for item in cli_result.ranked_opportunities
    ]
    assert len(runner_order) == len(export_order)


def test_opportunity_ranking_preserves_profit_roi_and_decision(tmp_path: Path) -> None:
    runner = _runner()
    result = run_discovery_command(
        DiscoveryCommandOptions(
            brands=["Chanel"],
            keyword="wallet",
            max_results_per_brand=2,
            export=True,
        ),
        runner=runner,
    )

    baseline = {
        candidate.supplier_product.external_id: {
            "profit_jpy": candidate.profit_result.profit_jpy if candidate.profit_result else None,
            "roi": candidate.profit_result.roi if candidate.profit_result else None,
            "decision": candidate.buy_decision.decision if candidate.buy_decision else None,
        }
        for candidate in result.ranked_candidates
        if candidate.status is DiscoveryCandidateStatus.SUCCESS
    }

    for item in result.ranked_opportunities:
        external_id = item.candidate.supplier_product.external_id
        expected = baseline[external_id]
        profit_result = item.candidate.profit_result
        assert profit_result is not None
        assert profit_result.profit_jpy == expected["profit_jpy"]
        assert profit_result.roi == expected["roi"]
        assert item.candidate.buy_decision is not None
        assert item.candidate.buy_decision.decision == expected["decision"]

    export_path = export_discovery_report(result, output_dir=tmp_path)
    assert export_path is not None
