"""Integration tests for cross-marketplace comparison."""

from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

import openpyxl

from comparison.demo_providers import build_comparison_demo_marketplaces
from comparison.engine import ComparisonEngine
from comparison.models import MarketplaceCandidate
from comparison.service import CrossMarketplaceComparisonService
from config.constants import SHEET_MARKETPLACE_COMPARISON
from excel.exporter import ExcelExporter
from excel.template import MARKETPLACE_COMPARISON_COLUMNS
from main import build_phase3_calculator, build_phase3_products, run_comparison_demo
from models.marketplace_search_result import MarketplaceSearchResult
from models.price_result import CALCULATION_SUCCESS, CALCULATION_UNKNOWN_CURRENCY, PriceResult
from models.product import Product


def test_comparison_demo_export(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr("main.OUTPUT_DIR", tmp_path)
    monkeypatch.setattr("main.EXCEL_FILENAME", "comparison.xlsx")
    with patch("utils.http.fetch_url"), patch("utils.http.HttpClient"):
        output_path = run_comparison_demo()
    assert output_path.exists()
    workbook = openpyxl.load_workbook(output_path)
    assert SHEET_MARKETPLACE_COMPARISON in workbook.sheetnames
    sheet = workbook[SHEET_MARKETPLACE_COMPARISON]
    headers = [sheet.cell(row=1, column=i + 1).value for i in range(sheet.max_column)]
    assert headers == MARKETPLACE_COMPARISON_COLUMNS
    assert sheet.max_row == 4
    assert "selected_review_marketplace" in headers
    assert "highest_profit_marketplace" in headers


def test_existing_sheets_unchanged(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr("main.OUTPUT_DIR", tmp_path)
    products = build_phase3_products()
    marketplaces = build_comparison_demo_marketplaces()
    run = CrossMarketplaceComparisonService().compare_products(
        products,
        marketplaces,
        build_phase3_calculator(),
    )
    exporter = ExcelExporter(output_dir=tmp_path, filename="cmp.xlsx")
    path = exporter.export_phase3_workbook(
        products,
        run.all_listings,
        run.ranked_price_results,
        comparison_results=run.products,
    )
    workbook = openpyxl.load_workbook(path)
    assert "Profit Ranking" in workbook.sheetnames
    assert "Domestic Listings" in workbook.sheetnames
    assert "Profit Analysis" in workbook.sheetnames
    assert "ROI Ranking" in workbook.sheetnames


def test_excel_non_jpy_candidate_representation(tmp_path: Path) -> None:
    product = Product(name="Demo", brand="Demo", model="M", sku="S")
    jpy = MarketplaceCandidate(
        marketplace_name="stockx",
        search_result=MarketplaceSearchResult(
            product=product,
            marketplace_name="stockx",
            selected_price_jpy=Decimal("10000"),
        ),
        price_result=PriceResult(
            product=product,
            profit_jpy=Decimal("1000"),
            profit_margin=Decimal("0.1"),
            calculation_status=CALCULATION_SUCCESS,
            domestic_sale_price_jpy=Decimal("10000"),
        ),
        listing_currency="JPY",
        source_price_amount=Decimal("10000"),
        jpy_comparable=True,
    )
    usd = MarketplaceCandidate(
        marketplace_name="goat",
        search_result=MarketplaceSearchResult(
            product=product,
            marketplace_name="goat",
            selected_price_jpy=Decimal("50000"),
        ),
        price_result=PriceResult(
            product=product,
            calculation_status=CALCULATION_UNKNOWN_CURRENCY,
            metadata={"source_currency": "USD", "source_price_amount": "50000"},
        ),
        listing_currency="USD",
        source_price_amount=Decimal("50000"),
        jpy_comparable=False,
        order_index=1,
    )
    comparison = ComparisonEngine().compare_product(product, [jpy, usd])
    exporter = ExcelExporter(output_dir=tmp_path, filename="mixed.xlsx")
    path = exporter.export_phase3_workbook(
        [product],
        [],
        [jpy.price_result],
        comparison_results=[comparison],
    )
    workbook = openpyxl.load_workbook(path)
    sheet = workbook[SHEET_MARKETPLACE_COMPARISON]
    row = {sheet.cell(row=1, column=i + 1).value: sheet.cell(row=2, column=i + 1).value for i in range(sheet.max_column)}
    assert row["selected_review_marketplace"] == "stockx"
    assert row["highest_profit_marketplace"] == "stockx"
    assert "USD" in str(row["currencies_observed"])
    assert row["selected_review_profit_jpy"] == 1000.0
    assert "Review" in str(row["recommendation"])
