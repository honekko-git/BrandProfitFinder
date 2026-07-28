"""Integration tests for Amazon marketplace factory and profit pipeline."""

import json
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from config.constants import MARKETPLACE_AMAZON_JP, SHEET_DOMESTIC_LISTINGS
from excel.exporter import ExcelExporter, SHEET_PRICE_RESULTS
from excel.template import MARKETPLACE_LISTING_COLUMNS
from marketplace.amazon_client import FakeAmazonClient
from marketplace.amazon_marketplace import AmazonMarketplace
from marketplace.amazon_settings import AmazonConfig
from marketplace.base_marketplace import BaseMarketplace
from marketplace.marketplace_factory import create_marketplace, get_all_marketplaces
from marketplace.local_marketplace import LocalMarketplace
from models.marketplace_listing import MarketplaceListing
from models.product import Product
from price_compare.marketplace_profit_service import calculate_profit_from_search_results
from price_compare.price_comparator import PriceComparator, PriceSelectionStrategy
from price_compare.profit_calculator import ProfitCalculator
from price_compare.profit_config import ProfitConfig

FIXTURES = Path(__file__).parent / "fixtures"


def _amazon_config(**overrides) -> AmazonConfig:
    defaults = dict(
        marketplace_id="A1VC38T7YXB528",
        default_currency="JPY",
        default_language="ja_JP",
        max_results=20,
        timeout_seconds=10,
        retry_count=0,
        enabled=True,
        demo_enabled=True,
    )
    defaults.update(overrides)
    return AmazonConfig(**defaults)


def _amazon_marketplace(name: str = "amazon_search_multiple.json") -> AmazonMarketplace:
    payload = json.loads((FIXTURES / name).read_text(encoding="utf-8"))
    return AmazonMarketplace(client=FakeAmazonClient(payload), config=_amazon_config())


def test_create_amazon_marketplace() -> None:
    marketplace = create_marketplace("amazon_jp", amazon_settings=_amazon_config())
    assert isinstance(marketplace, AmazonMarketplace)
    assert isinstance(marketplace, BaseMarketplace)


def test_create_amazon_alias() -> None:
    marketplace = create_marketplace("amazon", amazon_settings=_amazon_config())
    assert isinstance(marketplace, AmazonMarketplace)


def test_amazon_case_insensitive() -> None:
    marketplace = create_marketplace("AMAZON_JP", amazon_settings=_amazon_config())
    assert isinstance(marketplace, AmazonMarketplace)


def test_local_and_yahoo_still_work() -> None:
    assert isinstance(create_marketplace("local"), LocalMarketplace)
    assert isinstance(create_marketplace("yahoo"), BaseMarketplace)


def test_unknown_marketplace_raises() -> None:
    with pytest.raises(ValueError, match="Unsupported marketplace"):
        create_marketplace("ebay")


def test_get_all_marketplaces_includes_amazon() -> None:
    marketplaces = get_all_marketplaces(amazon_settings=_amazon_config())
    names = [item.marketplace_name for item in marketplaces]
    assert "local" in names
    assert "yahoo" in names
    assert MARKETPLACE_AMAZON_JP in names
    assert "rakuten" in names
    assert len(marketplaces) == 5


def test_profit_service_with_amazon_price() -> None:
    marketplace = _amazon_marketplace("amazon_search_normal.json")
    product = Product(name="Gucci Wallet", brand="GUCCI", model="456126", price=600.0, currency="USD", exchange_rate=150.0)
    result = marketplace.search(product)
    calculator = ProfitCalculator(
        ProfitConfig(
            international_shipping_jpy=Decimal("2500"),
            customs_duty_rate=Decimal("0.08"),
            import_tax_rate=Decimal("0.10"),
            domestic_shipping_jpy=Decimal("800"),
            marketplace_fee_rate=Decimal("0.12"),
            other_costs_jpy=Decimal("500"),
        )
    )
    price_result = calculate_profit_from_search_results([result], calculator)[0]
    assert price_result.domestic_sale_price_jpy == Decimal("128000")
    assert price_result.profit_jpy is not None


def test_profit_does_not_auto_deduct_points() -> None:
    marketplace = _amazon_marketplace("amazon_search_normal.json")
    product = Product(name="Gucci Wallet", brand="GUCCI", model="456126", price=600.0, currency="USD", exchange_rate=150.0)
    result = marketplace.search(product)
    listing = result.selected_listing
    assert listing is not None
    assert listing.points_jpy == Decimal("1280")
    calculator = ProfitCalculator()
    price_result = calculate_profit_from_search_results([result], calculator)[0]
    assert price_result.domestic_sale_price_jpy == Decimal("128000")


def test_profit_free_shipping() -> None:
    marketplace = _amazon_marketplace("amazon_search_normal.json")
    product = Product(name="Gucci Wallet", brand="GUCCI", model="456126", price=600.0, currency="USD", exchange_rate=150.0)
    result = marketplace.search(product)
    assert result.selected_listing is not None
    assert result.selected_listing.shipping_jpy == Decimal("0")
    assert result.selected_price_jpy == Decimal("128000")


def test_profit_with_shipping() -> None:
    marketplace = _amazon_marketplace("amazon_search_multiple.json")
    product = Product(name="Gucci Bag", brand="GUCCI", model="GG-MARMONT", price=600.0, currency="USD", exchange_rate=150.0)
    result = marketplace.search(product)
    selected = PriceComparator().select_from_listings(result.valid_listings, strategy=PriceSelectionStrategy.LOWEST)
    assert selected is not None
    assert selected[1] == Decimal("88000")


def test_profit_unknown_shipping_uses_item_price_only() -> None:
    payload = {
        "items": [
            {
                "asin": "B0UNK001",
                "title": "Unknown Shipping Item",
                "brand": "GUCCI",
                "model_number": "456126",
                "price": {"amount": 100000, "currency": "JPY"},
                "detail_page_url": "https://www.amazon.co.jp/dp/B0UNK001",
            }
        ]
    }
    marketplace = AmazonMarketplace(client=FakeAmazonClient(payload), config=_amazon_config())
    product = Product(name="Gucci Wallet", brand="GUCCI", model="456126", price=600.0, currency="USD", exchange_rate=150.0)
    result = marketplace.search(product)
    assert result.selected_price_jpy == Decimal("100000")


def test_no_auto_amazon_fee_added() -> None:
    calculator = ProfitCalculator(ProfitConfig(marketplace_fee_rate=Decimal("0")))
    marketplace = _amazon_marketplace("amazon_search_normal.json")
    product = Product(name="Gucci Wallet", brand="GUCCI", model="456126", price=600.0, currency="USD", exchange_rate=150.0)
    result = marketplace.search(product)
    price_result = calculate_profit_from_search_results([result], calculator)[0]
    assert price_result.marketplace_fee_jpy == Decimal("0")


def test_excel_amazon_rows(tmp_path: Path) -> None:
    marketplace = _amazon_marketplace("amazon_search_multiple.json")
    product = Product(name="Gucci Bag", brand="GUCCI", model="GG-MARMONT", price=600.0, currency="USD", exchange_rate=150.0)
    result = marketplace.search(product)
    calculator = ProfitCalculator()
    price_results = calculate_profit_from_search_results([result], calculator)
    exporter = ExcelExporter(output_dir=tmp_path, filename="amazon.xlsx")
    output_path = exporter.export_phase3_workbook([product], result.listings, price_results)

    workbook = openpyxl.load_workbook(output_path)
    sheet = workbook[SHEET_DOMESTIC_LISTINGS]
    marketplace_col = MARKETPLACE_LISTING_COLUMNS.index("marketplace_name") + 1
    values = [sheet.cell(row=row, column=marketplace_col).value for row in range(2, sheet.max_row + 1)]
    assert MARKETPLACE_AMAZON_JP in values

    price_col = MARKETPLACE_LISTING_COLUMNS.index("price_jpy") + 1
    price_value = sheet.cell(row=2, column=price_col).value
    assert isinstance(price_value, (int, float))

    points_col = MARKETPLACE_LISTING_COLUMNS.index("points_jpy") + 1
    points_value = sheet.cell(row=2, column=points_col).value
    assert points_value is None or isinstance(points_value, (int, float))
    assert points_value != "None"

    url_col = MARKETPLACE_LISTING_COLUMNS.index("listing_url") + 1
    url_cell = sheet.cell(row=2, column=url_col)
    assert url_cell.hyperlink is not None


def test_excel_empty_amazon_results(tmp_path: Path) -> None:
    marketplace = _amazon_marketplace("amazon_search_empty.json")
    product = Product(name="Unknown", brand="Demo", price=100.0)
    result = marketplace.search(product)
    exporter = ExcelExporter(output_dir=tmp_path, filename="amazon_empty.xlsx")
    output_path = exporter.export_phase3_workbook([product], result.listings, [])
    workbook = openpyxl.load_workbook(output_path)
    assert SHEET_DOMESTIC_LISTINGS in workbook.sheetnames
    assert SHEET_PRICE_RESULTS in workbook.sheetnames


def test_excel_yahoo_columns_preserved(tmp_path: Path) -> None:
    yahoo_listing = MarketplaceListing(
        marketplace_name="yahoo",
        listing_id="Y-1",
        title="Yahoo Item",
        price_jpy=Decimal("10000"),
        shipping_jpy=Decimal("0"),
        listing_url="https://example.com/yahoo",
    )
    amazon_listing = _amazon_marketplace("amazon_search_normal.json").search(
        Product(name="Gucci Wallet", brand="GUCCI", model="456126", price=100.0)
    ).listings[0]
    exporter = ExcelExporter(output_dir=tmp_path, filename="mixed.xlsx")
    exporter.export_phase3_workbook([], [yahoo_listing, amazon_listing], [])
    workbook = openpyxl.load_workbook(exporter.output_path)
    sheet = workbook[SHEET_DOMESTIC_LISTINGS]
    assert sheet.max_row == 3
