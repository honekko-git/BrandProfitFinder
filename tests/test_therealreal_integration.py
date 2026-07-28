"""Integration tests for The RealReal marketplace."""

import json
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from config.constants import MARKETPLACE_THEREALREAL, SHEET_DOMESTIC_LISTINGS
from excel.exporter import ExcelExporter
from excel.template import MARKETPLACE_LISTING_COLUMNS
from marketplace.marketplace_factory import create_marketplace
from marketplace.therealreal_client import FakeTheRealRealClient
from marketplace.therealreal_exceptions import TheRealRealConfigurationError
from marketplace.therealreal_marketplace import TheRealRealMarketplace, create_therealreal_marketplace
from marketplace.therealreal_settings import TheRealRealSettings
from models.product import Product
from price_compare.marketplace_profit_service import calculate_profit_from_search_results
from price_compare.profit_calculator import ProfitCalculator

FIXTURES = Path(__file__).parent / "fixtures"


def _settings(**overrides) -> TheRealRealSettings:
    defaults = dict(
        enabled=True,
        timeout_seconds=10,
        page_size=20,
        max_pages=2,
        default_currency="JPY",
        demo_fixture_path="therealreal_search_normal.json",
        allow_unknown_currency=False,
        include_unavailable=False,
        include_sold=False,
        include_reserved=False,
        include_final_sale=True,
        include_discounted_only=False,
    )
    defaults.update(overrides)
    return TheRealRealSettings(**defaults)


def _marketplace(name: str = "therealreal_search_normal.json", **settings) -> TheRealRealMarketplace:
    payload = json.loads((FIXTURES / name).read_text(encoding="utf-8"))
    return TheRealRealMarketplace(
        client=FakeTheRealRealClient(payload),
        settings=_settings(**settings),
    )


def test_marketplace_search() -> None:
    product = Product(
        name="Marmont Bag",
        brand="GUCCI",
        model="GG-MARMONT",
        price=600.0,
        currency="USD",
        exchange_rate=150.0,
    )
    result = _marketplace().search(product)
    assert result.marketplace_name == MARKETPLACE_THEREALREAL
    assert len(result.listings) == 1
    assert result.selected_price_jpy == Decimal("168000")


def test_empty_query() -> None:
    result = _marketplace().search(Product(name="", brand="", model=""), query="  ")
    assert result.status == "error"


def test_pagination_two_pages() -> None:
    p1 = json.loads((FIXTURES / "therealreal_search_page_1.json").read_text(encoding="utf-8"))
    p2 = json.loads((FIXTURES / "therealreal_search_page_2.json").read_text(encoding="utf-8"))
    mp = TheRealRealMarketplace(
        client=FakeTheRealRealClient(pages={1: p1, 2: p2}),
        settings=_settings(max_pages=2),
    )
    result = mp.search(Product(name="Bag", brand="GUCCI", model="GG-MARMONT", price=100.0))
    assert len(result.listings) == 2


def test_factory_aliases() -> None:
    payload = json.loads((FIXTURES / "therealreal_search_normal.json").read_text(encoding="utf-8"))
    client = FakeTheRealRealClient(payload)
    settings = _settings()
    for alias in ("the_real_real", "the-real-real", "realreal", "trr"):
        mp = create_marketplace(alias, therealreal_client=client, therealreal_settings=settings)
        assert isinstance(mp, TheRealRealMarketplace)


def test_factory_no_fake_without_client() -> None:
    with pytest.raises(TheRealRealConfigurationError):
        create_marketplace("therealreal", therealreal_settings=_settings())


def test_profit_metadata() -> None:
    product = Product(name="Bag", brand="GUCCI", model="GG-MARMONT", price=600.0, currency="USD", exchange_rate=150.0)
    result = _marketplace().search(product)
    price_result = calculate_profit_from_search_results([result], ProfitCalculator())[0]
    assert price_result.domestic_sale_price_jpy == Decimal("168000")
    assert price_result.metadata.get("discount_applied") is False
    assert price_result.metadata.get("final_sale") is False


def test_excel_rows(tmp_path: Path) -> None:
    product = Product(name="Bag", brand="GUCCI", model="GG-MARMONT", price=600.0, currency="USD", exchange_rate=150.0)
    result = _marketplace("therealreal_search_multiple.json").search(product)
    price_results = calculate_profit_from_search_results([result], ProfitCalculator())
    output_path = ExcelExporter(output_dir=tmp_path, filename="trr.xlsx").export_phase3_workbook(
        [product], result.listings, price_results
    )
    sheet = openpyxl.load_workbook(output_path)[SHEET_DOMESTIC_LISTINGS]
    col = MARKETPLACE_LISTING_COLUMNS.index("marketplace_name") + 1
    assert MARKETPLACE_THEREALREAL in [sheet.cell(row=r, column=col).value for r in range(2, sheet.max_row + 1)]
