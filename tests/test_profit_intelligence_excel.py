"""Tests for profit intelligence Excel export."""

from decimal import Decimal
from pathlib import Path

import openpyxl

from excel.exporter import ExcelExporter, SHEET_PRICE_RESULTS
from excel.template import PRICE_RESULT_COLUMNS, PROFIT_INTELLIGENCE_COLUMNS, price_result_columns
from models.price_result import PriceResult
from models.product import Product
from profit_intelligence.models import ProfitIntelligenceResult
from profit_intelligence.service import ProfitIntelligenceService
from price_compare.profit_calculator import ProfitCalculator


def test_intelligence_columns_appended_when_scored(tmp_path: Path) -> None:
    product = Product(name="Bag", brand="Demo", price=100.0, currency="USD", exchange_rate=150.0)
    result = ProfitCalculator().calculate(product, Decimal("50000"))
    service = ProfitIntelligenceService()
    scored = service.score_results([result])
    assert price_result_columns(scored) == PRICE_RESULT_COLUMNS + PROFIT_INTELLIGENCE_COLUMNS

    exporter = ExcelExporter(output_dir=tmp_path, filename="intel.xlsx")
    output_path = exporter.export_phase3_workbook([product], [], scored)
    workbook = openpyxl.load_workbook(output_path)
    sheet = workbook[SHEET_PRICE_RESULTS]
    headers = [sheet.cell(row=1, column=i + 1).value for i in range(sheet.max_column)]
    assert headers == price_result_columns(scored)


def test_legacy_export_without_intelligence_columns(tmp_path: Path) -> None:
    product = Product(name="Bag", brand="Demo", price=100.0, currency="USD", exchange_rate=150.0)
    result = ProfitCalculator().calculate(product, Decimal("50000"))
    exporter = ExcelExporter(output_dir=tmp_path, filename="legacy.xlsx")
    exporter.export_phase3_workbook([product], [], [result])
    workbook = openpyxl.load_workbook(exporter.output_path)
    sheet = workbook[SHEET_PRICE_RESULTS]
    headers = [sheet.cell(row=1, column=i + 1).value for i in range(len(PRICE_RESULT_COLUMNS))]
    assert headers == PRICE_RESULT_COLUMNS


def test_none_scores_not_written_as_zero(tmp_path: Path) -> None:
    product = Product(name="Bag", brand="Demo", price=100.0, currency="USD", exchange_rate=150.0)
    result = ProfitCalculator().calculate(product, Decimal("50000"))
    result.profit_intelligence = ProfitIntelligenceResult(
        overall_score=None,
        profit_score=None,
        velocity_score=None,
        risk_score=None,
        confidence_score=10.0,
        recommendation="Insufficient data",
        recommendation_stars=0,
        scoring_version="profit-intelligence-v1",
    )
    exporter = ExcelExporter(output_dir=tmp_path, filename="none_scores.xlsx")
    exporter.export_phase3_workbook([product], [], [result])
    workbook = openpyxl.load_workbook(exporter.output_path)
    sheet = workbook[SHEET_PRICE_RESULTS]
    columns = price_result_columns([result])
    overall_col = columns.index("overall_score") + 1
    assert sheet.cell(row=2, column=overall_col).value in (None, "")


def test_legacy_column_order_preserved_when_intelligence_enabled(tmp_path: Path) -> None:
    product = Product(name="Bag", brand="Demo", price=100.0, currency="USD", exchange_rate=150.0)
    result = ProfitCalculator().calculate(product, Decimal("50000"))
    scored = ProfitIntelligenceService().score_results([result])
    exporter = ExcelExporter(output_dir=tmp_path, filename="order.xlsx")
    exporter.export_phase3_workbook([product], [], scored)
    workbook = openpyxl.load_workbook(exporter.output_path)
    sheet = workbook[SHEET_PRICE_RESULTS]
    headers = [sheet.cell(row=1, column=i + 1).value for i in range(sheet.max_column)]
    assert headers[: len(PRICE_RESULT_COLUMNS)] == PRICE_RESULT_COLUMNS
    assert headers[len(PRICE_RESULT_COLUMNS) :] == PROFIT_INTELLIGENCE_COLUMNS


def test_numeric_zero_score_exports_as_zero_not_blank(tmp_path: Path) -> None:
    product = Product(name="Bag", brand="Demo", price=100.0, currency="USD", exchange_rate=150.0)
    result = ProfitCalculator().calculate(product, Decimal("50000"))
    result.profit_intelligence = ProfitIntelligenceResult(
        overall_score=0.0,
        profit_score=0.0,
        velocity_score=0.0,
        risk_score=0.0,
        confidence_score=0.0,
        recommendation="Low-priority review candidate",
        recommendation_stars=1,
        scoring_version="profit-intelligence-v1",
    )
    exporter = ExcelExporter(output_dir=tmp_path, filename="zero.xlsx")
    exporter.export_phase3_workbook([product], [], [result])
    workbook = openpyxl.load_workbook(exporter.output_path)
    sheet = workbook[SHEET_PRICE_RESULTS]
    columns = price_result_columns([result])
    overall_col = columns.index("overall_score") + 1
    assert sheet.cell(row=2, column=overall_col).value == 0.0


def test_reason_warning_and_missing_delimiters(tmp_path: Path) -> None:
    product = Product(name="Bag", brand="Demo", price=100.0, currency="USD", exchange_rate=150.0)
    result = ProfitCalculator().calculate(product, Decimal("50000"))
    result.profit_intelligence = ProfitIntelligenceResult(
        overall_score=55.0,
        profit_score=60.0,
        velocity_score=None,
        risk_score=20.0,
        confidence_score=40.0,
        recommendation="Cautious review candidate",
        recommendation_stars=2,
        reasons=("Reason one", "Reason two"),
        warnings=("Warning one",),
        missing_fields=("sales_last_30_days", "asks_count"),
        scoring_version="profit-intelligence-v1",
    )
    payload = result.to_dict()
    assert payload["score_reasons"] == "Reason one | Reason two"
    assert payload["score_warnings"] == "Warning one"
    assert payload["missing_score_data"] == "sales_last_30_days | asks_count"
    assert payload["scoring_version"] == "profit-intelligence-v1"


def test_workbook_formatting_preserved_with_intelligence_columns(tmp_path: Path) -> None:
    product = Product(name="Bag", brand="Demo", price=100.0, currency="USD", exchange_rate=150.0)
    result = ProfitCalculator().calculate(product, Decimal("50000"))
    scored = ProfitIntelligenceService().score_results([result])
    exporter = ExcelExporter(output_dir=tmp_path, filename="format.xlsx")
    exporter.export_phase3_workbook([product], [], scored)
    workbook = openpyxl.load_workbook(exporter.output_path)
    sheet = workbook[SHEET_PRICE_RESULTS]
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref is not None
