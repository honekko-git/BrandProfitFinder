"""StockX demo CLI regression tests for Profit Intelligence integration."""

from __future__ import annotations

import logging
import sys
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest

import main
from excel.exporter import SHEET_PRICE_RESULTS
from excel.template import PRICE_RESULT_COLUMNS, PROFIT_INTELLIGENCE_COLUMNS
from profit_intelligence.service import ProfitIntelligenceService


@pytest.fixture
def stockx_output_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    monkeypatch.setattr(main, "OUTPUT_DIR", tmp_path)
    return tmp_path


def test_stockx_demo_without_intelligence_skips_scoring(stockx_output_dir: Path) -> None:
    with patch("utils.http.fetch_url"), patch("utils.http.HttpClient"):
        with patch.object(sys, "argv", ["main.py", "--demo-stockx"]):
            with patch.object(ProfitIntelligenceService, "score_results") as score_mock:
                main.main()
                score_mock.assert_not_called()

    workbook = openpyxl.load_workbook(stockx_output_dir / main.EXCEL_FILENAME)
    sheet = workbook[SHEET_PRICE_RESULTS]
    assert sheet.max_column == len(PRICE_RESULT_COLUMNS)
    headers = [sheet.cell(row=1, column=i + 1).value for i in range(sheet.max_column)]
    assert headers == PRICE_RESULT_COLUMNS


def test_stockx_demo_with_intelligence_scores_and_exports(
    stockx_output_dir: Path,
    caplog: pytest.LogCaptureFixture,
) -> None:
    captured_profits: list[tuple[Decimal, Decimal]] = []
    original_score_results = ProfitIntelligenceService.score_results

    def recording_score_results(
        self: ProfitIntelligenceService,
        results,
        search_results=None,
    ):
        captured_profits.extend([(item.profit_jpy, item.profit_margin) for item in results])
        scored = original_score_results(self, results, search_results)
        for before, after in zip(captured_profits, scored, strict=True):
            assert after.profit_jpy == before[0]
            assert after.profit_margin == before[1]
        assert len(scored) == 3
        assert all(item.profit_intelligence is not None for item in scored)
        assert all(item.profit_intelligence.overall_score is not None for item in scored)
        return scored

    caplog.set_level(logging.INFO)
    with patch("utils.http.fetch_url"), patch("utils.http.HttpClient"):
        with patch.object(sys, "argv", ["main.py", "--demo-stockx", "--profit-intelligence"]):
            with patch.object(
                ProfitIntelligenceService,
                "score_results",
                recording_score_results,
            ):
                main.main()

    summary_messages = [
        record.message
        for record in caplog.records
        if record.name == "profit_intelligence.service"
    ]
    assert any(
        "Profit intelligence summary: scored=3 complete=3 insufficient=0" in message
        for message in summary_messages
    )

    workbook = openpyxl.load_workbook(stockx_output_dir / main.EXCEL_FILENAME)
    sheet = workbook[SHEET_PRICE_RESULTS]
    headers = [sheet.cell(row=1, column=i + 1).value for i in range(sheet.max_column)]
    assert headers == PRICE_RESULT_COLUMNS + PROFIT_INTELLIGENCE_COLUMNS
    assert sheet.max_row == 4  # header + 3 scored StockX results

    scoring_version_col = headers.index("scoring_version") + 1
    assert sheet.cell(row=2, column=scoring_version_col).value == "profit-intelligence-v1"
