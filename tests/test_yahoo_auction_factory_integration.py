"""Integration tests for Yahoo Auction marketplace factory and profit pipeline."""

import json
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from config.constants import MARKETPLACE_YAHOO_AUCTION, SHEET_DOMESTIC_LISTINGS
from excel.exporter import ExcelExporter
from excel.template import MARKETPLACE_LISTING_COLUMNS
from marketplace.base_marketplace import BaseMarketplace
from marketplace.marketplace_factory import create_marketplace, get_all_marketplaces
from marketplace.local_marketplace import LocalMarketplace
from marketplace.yahoo_auction_client import FakeYahooAuctionClient
from marketplace.yahoo_auction_marketplace import YahooAuctionMarketplace
from marketplace.yahoo_auction_settings import YahooAuctionConfig
from models.marketplace_listing import MarketplaceListing
from models.product import Product
from price_compare.marketplace_profit_service import calculate_profit_from_search_results
from price_compare.price_comparator import PriceComparator, PriceSelectionStrategy
from price_compare.profit_calculator import ProfitCalculator
from price_compare.profit_config import ProfitConfig

FIXTURES = Path(__file__).parent / "fixtures"


def _config(**overrides) -> YahooAuctionConfig:
    defaults = dict(
        enabled=True,
        demo_enabled=True,
        data_source="",
        timeout_seconds=10,
        max_retries=0,
        hits=20,
        sort="end_time",
    )
    defaults.update(overrides)
    return YahooAuctionConfig(**defaults)


def _marketplace(name: str = "yahoo_auction_search_multiple.json") -> YahooAuctionMarketplace:
    payload = json.loads((FIXTURES / name).read_text(encoding="utf-8"))
    return YahooAuctionMarketplace(client=FakeYahooAuctionClient(payload), config=_config())


def test_create_yahoo_auction_marketplace() -> None:
    marketplace = create_marketplace("yahoo_auction", yahoo_auction_settings=_config())
    assert isinstance(marketplace, YahooAuctionMarketplace)
    assert isinstance(marketplace, BaseMarketplace)


def test_create_yahoo_auction_alias() -> None:
    marketplace = create_marketplace("yahoo-auction", yahoo_auction_settings=_config())
    assert isinstance(marketplace, YahooAuctionMarketplace)


def test_yahooauction_alias() -> None:
    marketplace = create_marketplace("yahooauction", yahoo_auction_settings=_config())
    assert isinstance(marketplace, YahooAuctionMarketplace)


def test_auctions_alias() -> None:
    marketplace = create_marketplace("auctions", yahoo_auction_settings=_config())
    assert isinstance(marketplace, YahooAuctionMarketplace)


def test_local_and_other_marketplaces_preserved() -> None:
    assert isinstance(create_marketplace("local"), LocalMarketplace)
    assert create_marketplace("yahoo").marketplace_name == "yahoo"
    assert create_marketplace("amazon_jp").marketplace_name == "amazon_jp"
    assert create_marketplace("rakuten").marketplace_name == "rakuten"


def test_unknown_marketplace_raises() -> None:
    with pytest.raises(ValueError, match="Unsupported marketplace"):
        create_marketplace("ebay")


def test_get_all_marketplaces_includes_yahoo_auction() -> None:
    marketplaces = get_all_marketplaces(yahoo_auction_settings=_config())
    names = [item.marketplace_name for item in marketplaces]
    assert MARKETPLACE_YAHOO_AUCTION in names
    assert len(marketplaces) == 5


def test_profit_current_price() -> None:
    marketplace = _marketplace("yahoo_auction_search_active.json")
    product = Product(name="Wallet", brand="GUCCI", model="428726", price=600.0, currency="USD", exchange_rate=150.0)
    result = marketplace.search(product)
    selected = next(item for item in result.valid_listings if item.listing_id == "test-auction-active-001")
    assert selected.price_jpy == Decimal("88000")
    assert selected.source_metadata["price_source"] == "current_price"


def test_profit_buy_now_price() -> None:
    marketplace = _marketplace("yahoo_auction_search_normal.json")
    product = Product(name="Wallet", brand="GUCCI", model="428726", price=600.0, currency="USD", exchange_rate=150.0)
    result = marketplace.search(product)
    assert result.selected_price_jpy == Decimal("108000")


def test_profit_winning_price() -> None:
    marketplace = _marketplace("yahoo_auction_search_sold.json")
    product = Product(name="Wallet", brand="GUCCI", model="428726", price=600.0, currency="USD", exchange_rate=150.0)
    result = marketplace.search(product)
    sold = next(item for item in result.valid_listings if item.listing_id == "test-auction-sold-001")
    assert sold.price_jpy == Decimal("95000")


def test_profit_free_shipping() -> None:
    marketplace = _marketplace("yahoo_auction_search_multiple.json")
    product = Product(name="Wallet", brand="GUCCI", model="428726", price=600.0, currency="USD", exchange_rate=150.0)
    result = marketplace.search(product)
    free = next(item for item in result.valid_listings if item.listing_id == "test-auction-101")
    assert free.shipping_jpy == Decimal("0")
    selected = PriceComparator().select_from_listings([free], strategy=PriceSelectionStrategy.LOWEST)
    assert selected is not None
    assert selected[1] == Decimal("108000")


def test_profit_with_shipping() -> None:
    marketplace = _marketplace("yahoo_auction_search_multiple.json")
    product = Product(name="Bag", brand="GUCCI", model="GG-MARMONT", price=600.0, currency="USD", exchange_rate=150.0)
    result = marketplace.search(product)
    paid = next(item for item in result.valid_listings if item.listing_id == "test-auction-102")
    assert paid.shipping_jpy == Decimal("800")
    selected = PriceComparator().select_from_listings([paid], strategy=PriceSelectionStrategy.LOWEST)
    assert selected is not None
    assert selected[1] == Decimal("185800")


def test_profit_unknown_shipping_uses_item_price_only() -> None:
    marketplace = _marketplace("yahoo_auction_search_normal.json")
    product = Product(name="Wallet", brand="GUCCI", model="428726", price=600.0, currency="USD", exchange_rate=150.0)
    result = marketplace.search(product)
    assert result.selected_listing is not None
    assert result.selected_listing.shipping_unknown is True
    assert result.selected_price_jpy == Decimal("108000")


def test_no_auto_fee_added() -> None:
    calculator = ProfitCalculator(ProfitConfig(marketplace_fee_rate=Decimal("0")))
    marketplace = _marketplace("yahoo_auction_search_normal.json")
    product = Product(name="Wallet", brand="GUCCI", model="428726", price=600.0, currency="USD", exchange_rate=150.0)
    result = marketplace.search(product)
    price_result = calculate_profit_from_search_results([result], calculator)[0]
    assert price_result.marketplace_fee_jpy == Decimal("0")


def test_no_auto_points_reflected() -> None:
    marketplace = _marketplace("yahoo_auction_search_normal.json")
    product = Product(name="Wallet", brand="GUCCI", model="428726", price=600.0, currency="USD", exchange_rate=150.0)
    result = marketplace.search(product)
    listing = result.selected_listing
    assert listing is not None
    assert listing.points_jpy is None
    price_result = calculate_profit_from_search_results([result], ProfitCalculator())[0]
    assert price_result.domestic_sale_price_jpy == Decimal("108000")


def test_excel_yahoo_auction_rows(tmp_path: Path) -> None:
    marketplace = _marketplace("yahoo_auction_search_multiple.json")
    product = Product(name="Wallet", brand="GUCCI", model="428726", price=600.0, currency="USD", exchange_rate=150.0)
    result = marketplace.search(product)
    calculator = ProfitCalculator()
    price_results = calculate_profit_from_search_results([result], calculator)
    exporter = ExcelExporter(output_dir=tmp_path, filename="yahoo_auction.xlsx")
    output_path = exporter.export_phase3_workbook([product], result.listings, price_results)

    workbook = openpyxl.load_workbook(output_path)
    sheet = workbook[SHEET_DOMESTIC_LISTINGS]
    marketplace_col = MARKETPLACE_LISTING_COLUMNS.index("marketplace_name") + 1
    values = [sheet.cell(row=row, column=marketplace_col).value for row in range(2, sheet.max_row + 1)]
    assert MARKETPLACE_YAHOO_AUCTION in values

    price_col = MARKETPLACE_LISTING_COLUMNS.index("price_jpy") + 1
    price_value = sheet.cell(row=2, column=price_col).value
    assert isinstance(price_value, (int, float))

    price_source_col = MARKETPLACE_LISTING_COLUMNS.index("price_source") + 1
    price_source_value = sheet.cell(row=2, column=price_source_col).value
    assert price_source_value in {"buy_now_price", "current_price", "winning_price"}

    url_col = MARKETPLACE_LISTING_COLUMNS.index("listing_url") + 1
    url_cell = sheet.cell(row=2, column=url_col)
    assert url_cell.hyperlink is not None

    auction_id_col = MARKETPLACE_LISTING_COLUMNS.index("listing_id") + 1
    auction_id_value = sheet.cell(row=2, column=auction_id_col).value
    assert isinstance(auction_id_value, str)


def test_excel_none_not_string(tmp_path: Path) -> None:
    listing = MarketplaceListing(
        marketplace_name=MARKETPLACE_YAHOO_AUCTION,
        listing_id="test-auction-x",
        title="Test",
        price_jpy=Decimal("10000"),
        shipping_jpy=None,
        shipping_unknown=True,
        listing_url="https://example.invalid/auction/test-auction-x",
        source_metadata={"price_source": "current_price", "listing_status": "active"},
    )
    exporter = ExcelExporter(output_dir=tmp_path, filename="yahoo_auction_none.xlsx")
    exporter.export_phase3_workbook([], [listing], [])
    workbook = openpyxl.load_workbook(exporter.output_path)
    sheet = workbook[SHEET_DOMESTIC_LISTINGS]
    shipping_col = MARKETPLACE_LISTING_COLUMNS.index("shipping_jpy") + 1
    assert sheet.cell(row=2, column=shipping_col).value is None


def test_excel_empty_results(tmp_path: Path) -> None:
    marketplace = _marketplace("yahoo_auction_search_empty.json")
    product = Product(name="Unknown", brand="Demo", price=100.0)
    result = marketplace.search(product)
    exporter = ExcelExporter(output_dir=tmp_path, filename="yahoo_auction_empty.xlsx")
    output_path = exporter.export_phase3_workbook([product], result.listings, [])
    workbook = openpyxl.load_workbook(output_path)
    assert SHEET_DOMESTIC_LISTINGS in workbook.sheetnames


def test_excel_existing_marketplaces_preserved(tmp_path: Path) -> None:
    yahoo_listing = MarketplaceListing(
        marketplace_name="yahoo",
        listing_id="Y-1",
        title="Yahoo Item",
        price_jpy=Decimal("10000"),
        shipping_jpy=Decimal("0"),
        listing_url="https://example.com/yahoo",
    )
    auction_listing = _marketplace("yahoo_auction_search_normal.json").search(
        Product(name="Wallet", brand="GUCCI", model="428726", price=100.0)
    ).listings[0]
    exporter = ExcelExporter(output_dir=tmp_path, filename="mixed.xlsx")
    exporter.export_phase3_workbook([], [yahoo_listing, auction_listing], [])
    workbook = openpyxl.load_workbook(exporter.output_path)
    sheet = workbook[SHEET_DOMESTIC_LISTINGS]
    assert sheet.max_row == 3
