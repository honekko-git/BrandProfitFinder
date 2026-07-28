"""
Excel formatting helpers.
"""

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center")
PROFIT_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
LOSS_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
HYPERLINK_FONT = Font(color="0563C1", underline="single")


def format_header_row(sheet: Worksheet, column_count: int) -> None:
    """
    Apply header styling to the first row.

    Args:
        sheet: Target worksheet.
        column_count: Number of columns to style.
    """
    for col in range(1, column_count + 1):
        cell = sheet.cell(row=1, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT


def auto_fit_columns(sheet: Worksheet, max_width: int = 50) -> None:
    """
    Adjust column widths based on cell content.

    Args:
        sheet: Target worksheet.
        max_width: Maximum column width cap.
    """
    for column_cells in sheet.columns:
        length = max(len(str(cell.value or "")) for cell in column_cells)
        adjusted = min(max(length + 2, 10), max_width)
        sheet.column_dimensions[column_cells[0].column_letter].width = adjusted


def apply_sheet_layout(sheet: Worksheet, column_count: int) -> None:
    """
    Apply common sheet layout options.

    Args:
        sheet: Target worksheet.
        column_count: Number of populated columns.
    """
    format_header_row(sheet, column_count)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(column_count)}1"
    auto_fit_columns(sheet)


def apply_price_result_formatting(
    sheet: Worksheet,
    headers: list[str],
    row_count: int,
) -> None:
    """
    Apply profit/loss coloring and URL hyperlinks.

    Args:
        sheet: Target worksheet.
        headers: Header names in column order.
        row_count: Number of data rows excluding header.
    """
    header_index = {name: index + 1 for index, name in enumerate(headers)}
    profitable_col = header_index.get("is_profitable")
    url_cols = [header_index[name] for name in ("url", "image_url") if name in header_index]

    for row in range(2, row_count + 2):
        if profitable_col is not None:
            profitable_value = sheet.cell(row=row, column=profitable_col).value
            fill = PROFIT_FILL if profitable_value is True else LOSS_FILL
            for col in range(1, len(headers) + 1):
                sheet.cell(row=row, column=col).fill = fill

        for col in url_cols:
            cell = sheet.cell(row=row, column=col)
            if isinstance(cell.value, str) and cell.value.startswith(("http://", "https://")):
                cell.hyperlink = cell.value
                cell.font = HYPERLINK_FONT
