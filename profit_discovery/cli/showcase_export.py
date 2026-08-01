"""Showcase dashboard export helpers for discovery CLI reports."""

from __future__ import annotations

from decimal import Decimal

import pandas as pd
from openpyxl import load_workbook

from excel.formatter import apply_sheet_layout
from marketplace.domestic_market.execution import MarketExecutionResult
from profit_discovery.arbitrage.models import ArbitrageOpportunity
from profit_discovery.config.used_luxury import UsedLuxuryModeConfig
from profit_discovery.profit_ranking.models import UsedLuxuryProfitRankedResult
from profit_discovery.showcase.formatter import ShowcaseFormatter
from profit_discovery.showcase.models import ShowcaseOpportunity

SHEET_SHOWCASE = "Showcase"

SHOWCASE_COLUMNS: tuple[str, ...] = (
    "Rank",
    "Brand",
    "Product",
    "Supplier",
    "Profit",
    "ROI",
    "Demand Score",
    "Opportunity Score",
    "Decision",
    "Market Source",
    "Requested Mode",
    "Actual Source",
    "Fallback",
    "Business Mode",
    "Market Coverage",
    "Profit Rank",
    "Turnover Score",
    "Purchase Source",
    "Purchase URL",
    "Selling Market",
    "Selling URL",
    "Estimated Profit",
)


def showcase_opportunities_to_rows(opportunities: list[ShowcaseOpportunity]) -> list[dict[str, object]]:
    """Convert showcase opportunities into flat export rows."""
    rows: list[dict[str, object]] = []
    for item in opportunities:
        rows.append(
            {
                "Rank": item.rank,
                "Brand": item.brand,
                "Product": item.product_name,
                "Supplier": item.supplier,
                "Profit": _format_export_jpy(item.profit_jpy),
                "ROI": _format_export_percent(item.roi),
                "Demand Score": round(item.demand_score, 1),
                "Opportunity Score": round(item.opportunity_score, 1),
                "Decision": item.decision,
                "Market Source": item.market_source,
                "Requested Mode": item.requested_market_mode,
                "Actual Source": item.actual_market_source,
                "Fallback": _format_export_fallback(item.fallback_used),
                "Business Mode": item.business_mode,
                "Market Coverage": item.market_coverage,
                "Profit Rank": item.profit_rank if item.profit_rank is not None else "",
                "Turnover Score": (
                    round(item.turnover_score, 1) if item.turnover_score is not None else ""
                ),
                "Purchase Source": item.purchase_source,
                "Purchase URL": item.purchase_url,
                "Selling Market": item.selling_market,
                "Selling URL": item.selling_url,
                "Estimated Profit": _format_export_jpy(item.estimated_profit),
            }
        )
    return rows


def append_showcase_sheet(
    workbook_path,
    opportunities: list[ShowcaseOpportunity],
) -> None:
    """Append the Showcase sheet to an existing discovery workbook."""
    if not opportunities:
        return

    rows = showcase_opportunities_to_rows(opportunities)
    with pd.ExcelWriter(
        workbook_path,
        engine="openpyxl",
        mode="a",
        if_sheet_exists="replace",
    ) as writer:
        pd.DataFrame(rows, columns=list(SHOWCASE_COLUMNS)).to_excel(
            writer,
            index=False,
            sheet_name=SHEET_SHOWCASE,
        )

    workbook = load_workbook(workbook_path)
    if SHEET_SHOWCASE in workbook.sheetnames:
        apply_sheet_layout(workbook[SHEET_SHOWCASE], len(SHOWCASE_COLUMNS))
        workbook.save(workbook_path)


def build_showcase_opportunities_from_ranking(
    opportunities,
    *,
    market_execution: MarketExecutionResult | None = None,
    market_source: str | None = None,
    used_luxury_config: UsedLuxuryModeConfig | None = None,
    profit_ranking: list[UsedLuxuryProfitRankedResult] | None = None,
    arbitrage_ranking: list[ArbitrageOpportunity] | None = None,
) -> list[ShowcaseOpportunity]:
    """Build showcase rows from demand-integrated ranking results."""
    return list(
        ShowcaseFormatter().to_showcase_opportunities(
            opportunities,
            market_execution=market_execution,
            market_source=market_source,
            used_luxury_config=used_luxury_config,
            profit_ranking=profit_ranking,
            arbitrage_ranking=arbitrage_ranking,
        )
    )


def _format_export_jpy(value: Decimal | None) -> str:
    if value is None:
        return "N/A"
    return f"{int(value):,} JPY"


def _format_export_percent(value: Decimal | None) -> str:
    if value is None:
        return "N/A"
    return f"{value:.1f}%"


def _format_export_fallback(value: bool) -> str:
    return "YES" if value else "NO"
