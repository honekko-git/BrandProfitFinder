"""Demand-integrated opportunity ranking export helpers for discovery CLI reports."""

from __future__ import annotations

from decimal import Decimal

import pandas as pd
from openpyxl import load_workbook

from excel.formatter import apply_sheet_layout
from profit_discovery.opportunity.models import DemandIntegratedOpportunityResult

SHEET_DEMAND_OPPORTUNITY_RANKING = "Demand Opportunity Ranking"

DEMAND_OPPORTUNITY_RANKING_COLUMNS: tuple[str, ...] = (
    "Rank",
    "Product",
    "Supplier",
    "Profit",
    "Demand Score",
    "Opportunity Score",
    "ROI",
    "Decision",
)


def demand_opportunity_results_to_rows(
    opportunities: list[DemandIntegratedOpportunityResult],
) -> list[dict[str, object]]:
    """Convert ranked demand-integrated opportunity results into flat export rows."""
    rows: list[dict[str, object]] = []
    for item in opportunities:
        candidate = item.candidate
        profit_result = candidate.profit_result
        decision = candidate.buy_decision.decision.value if candidate.buy_decision else "N/A"
        product_name = (
            item.demand_profile.query
            if item.demand_profile is not None
            else candidate.supplier_product.title
        )
        rows.append(
            {
                "Rank": item.recommendation_rank,
                "Product": product_name,
                "Supplier": candidate.supplier_product.supplier_name,
                "Profit": _format_export_jpy(profit_result.profit_jpy if profit_result else None),
                "Demand Score": round(item.score.demand_score, 1),
                "Opportunity Score": round(item.score.total_score, 1),
                "ROI": _format_export_percent(profit_result.roi if profit_result else None),
                "Decision": decision,
            }
        )
    return rows


def append_demand_opportunity_ranking_sheet(
    workbook_path,
    opportunities: list[DemandIntegratedOpportunityResult],
) -> None:
    """Append the Demand Opportunity Ranking sheet to an existing discovery workbook."""
    if not opportunities:
        return

    rows = demand_opportunity_results_to_rows(opportunities)
    with pd.ExcelWriter(
        workbook_path,
        engine="openpyxl",
        mode="a",
        if_sheet_exists="replace",
    ) as writer:
        pd.DataFrame(rows, columns=list(DEMAND_OPPORTUNITY_RANKING_COLUMNS)).to_excel(
            writer,
            index=False,
            sheet_name=SHEET_DEMAND_OPPORTUNITY_RANKING,
        )

    workbook = load_workbook(workbook_path)
    if SHEET_DEMAND_OPPORTUNITY_RANKING in workbook.sheetnames:
        apply_sheet_layout(workbook[SHEET_DEMAND_OPPORTUNITY_RANKING], len(DEMAND_OPPORTUNITY_RANKING_COLUMNS))
        workbook.save(workbook_path)


def _format_export_jpy(value: Decimal | None) -> str:
    if value is None:
        return "N/A"
    return f"{int(value):,} JPY"


def _format_export_percent(value: Decimal | None) -> str:
    if value is None:
        return "N/A"
    return f"{value:.1f}%"
