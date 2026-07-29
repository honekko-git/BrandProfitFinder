"""Opportunity ranking export helpers for discovery CLI reports."""

from __future__ import annotations

from decimal import Decimal

import pandas as pd
from openpyxl import load_workbook

from excel.formatter import apply_sheet_layout
from profit_discovery.opportunity.models import OpportunityResult

SHEET_OPPORTUNITY_RANKING = "Opportunity Ranking"

OPPORTUNITY_RANKING_COLUMNS: tuple[str, ...] = (
    "Rank",
    "Product",
    "Supplier",
    "Score",
    "Profit",
    "Margin",
    "ROI",
    "Decision",
)


def opportunity_results_to_rows(opportunities: list[OpportunityResult]) -> list[dict[str, object]]:
    """Convert ranked opportunity results into flat export rows."""
    rows: list[dict[str, object]] = []
    for item in opportunities:
        candidate = item.candidate
        profit_result = candidate.profit_result
        decision = candidate.buy_decision.decision.value if candidate.buy_decision else "N/A"
        rows.append(
            {
                "Rank": item.recommendation_rank,
                "Product": candidate.supplier_product.title,
                "Supplier": candidate.supplier_product.supplier_name,
                "Score": round(item.score.total_score, 1),
                "Profit": _format_export_jpy(profit_result.profit_jpy if profit_result else None),
                "Margin": _format_export_percent(profit_result.profit_margin if profit_result else None),
                "ROI": _format_export_percent(profit_result.roi if profit_result else None),
                "Decision": decision,
            }
        )
    return rows


def append_opportunity_ranking_sheet(
    workbook_path,
    opportunities: list[OpportunityResult],
) -> None:
    """Append the Opportunity Ranking sheet to an existing discovery workbook."""
    if not opportunities:
        return

    rows = opportunity_results_to_rows(opportunities)
    with pd.ExcelWriter(
        workbook_path,
        engine="openpyxl",
        mode="a",
        if_sheet_exists="replace",
    ) as writer:
        pd.DataFrame(rows, columns=list(OPPORTUNITY_RANKING_COLUMNS)).to_excel(
            writer,
            index=False,
            sheet_name=SHEET_OPPORTUNITY_RANKING,
        )

    workbook = load_workbook(workbook_path)
    if SHEET_OPPORTUNITY_RANKING in workbook.sheetnames:
        apply_sheet_layout(workbook[SHEET_OPPORTUNITY_RANKING], len(OPPORTUNITY_RANKING_COLUMNS))
        workbook.save(workbook_path)


def _format_export_jpy(value: Decimal | None) -> str:
    if value is None:
        return "N/A"
    return f"{int(value):,} JPY"


def _format_export_percent(value: Decimal | None) -> str:
    if value is None:
        return "N/A"
    return f"{value:.1f}%"
