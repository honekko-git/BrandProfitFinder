"""Workbook validation helper for Phase 24 review."""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook

from config.constants import SHEET_MARKETPLACE_COMPARISON
from excel.template import MARKETPLACE_COMPARISON_COLUMNS

FORBIDDEN_SUBSTRINGS = (
    "IdentityDecision.",
    "EvidenceOutcome.",
    "probability",
    "authenticity",
    "authentic",
)


def main() -> None:
    wb = load_workbook("output/profit_ranking.xlsx")
    assert SHEET_MARKETPLACE_COMPARISON in wb.sheetnames, "Missing Marketplace Comparison sheet"
    sheet = wb[SHEET_MARKETPLACE_COMPARISON]
    headers = [sheet.cell(1, col).value for col in range(1, sheet.max_column + 1)]
    assert headers == MARKETPLACE_COMPARISON_COLUMNS, "Column order changed"
    rows = sheet.max_row - 1
    assert rows == 3, f"Expected 3 comparison rows, got {rows}"

    seen: set[tuple] = set()
    ranking_col = headers.index("ranking_score") + 1 if "ranking_score" in headers else None
    for row in range(2, sheet.max_row + 1):
        values = tuple(sheet.cell(row, col).value for col in range(1, sheet.max_column + 1))
        assert values not in seen, "Duplicate comparison row detected"
        seen.add(values)
        for value in values:
            if value is None:
                continue
            text = str(value)
            assert not text.startswith("<"), "Python repr leakage detected"
            lowered = text.lower()
            for forbidden in FORBIDDEN_SUBSTRINGS:
                assert forbidden not in lowered, f"Forbidden wording '{forbidden}' in cell: {text}"

    identity_cols = [name for name in headers if name.startswith("selected_review_identity_")]
    assert identity_cols, "Identity columns missing"

    print(
        f"workbook OK: {len(headers)} columns, {rows} rows, "
        f"{len(identity_cols)} identity columns, ranking_col={'yes' if ranking_col else 'no'}"
    )


if __name__ == "__main__":
    main()
